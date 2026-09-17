import docx, random, copy, re, os, csv, traceback, sys, json, threading, gc
from docx.document import Document
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT
import io

try:
    import openpyxl
    from openpyxl.styles import Font, Alignment
    from openpyxl.utils import get_column_letter
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

def iter_block_items(parent):
    if isinstance(parent, Document):
        parent_elm = parent.element.body
    else:
        parent_elm = parent._element
    for child in parent_elm.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, parent)
        elif isinstance(child, CT_Tbl):
            yield Table(child, parent)

def safe_append(body_elm, element):
    sectPr = body_elm.find(qn('w:sectPr'))
    if sectPr is not None:
        sectPr.addprevious(element)
    else:
        body_elm.append(element)

def get_resource_path(relative_path):
    try:
        base_path = sys._MEIPASS
    except Exception:
        base_path = os.path.dirname(os.path.abspath(__file__))
    return os.path.join(base_path, relative_path)

class DTMIXApp:
    def __init__(self, filepath, out_dir, so, truong, kythi, namhoc, monthi, thoigian, youngmix, ma_de_str):
        self.filepath = filepath
        self.output_dir = out_dir
        
        # Dữ liệu web truyền vào
        self.val_so_str = so
        self.val_truong_str = truong
        self.val_kythi_str = kythi
        self.val_namhoc_str = namhoc
        self.val_mon_str = monthi
        self.val_thoigian_str = thoigian
        self.ma_de_str = ma_de_str
        
        # Khởi tạo các biến nội tại gốc của bạn
        self.parsed_data = {"header": [], "parts": []}
        self.muc_entries = {}
        self.muc_shuffle_vars = {} 
        self.muc_fix_entries = {}  
        self.muc_order_shuffle_vars = {} 
        self.errors = []
        
        # Biến Dummy giả lập các biến Giao diện (Checkbox, Switch) cũ
        class DummyVar: 
            def __init__(self, val): self.val = val
            def get(self): return self.val
            
        self.youngmix_var = DummyVar(youngmix)
        self.continuous_numbering_var = DummyVar(False)
        self.master_fix_var = DummyVar(True)
        self.keep_muc_var = DummyVar(False)
        
        self.youngmix_data = [] 
        self.youngmix_total_q = 0
        self.ym_ui_controls = []

    class DummyWidget:
        def __init__(self, text): self.text = text
        def get(self): return self.text

    @property
    def val_so(self): return self.DummyWidget(self.val_so_str)
    @property
    def val_truong(self): return self.DummyWidget(self.val_truong_str)
    @property
    def val_kythi(self): return self.DummyWidget(self.val_kythi_str)
    @property
    def val_namhoc(self): return self.DummyWidget(self.val_namhoc_str)
    @property
    def val_mon(self): return self.DummyWidget(self.val_mon_str)
    @property
    def val_thoigian(self): return self.DummyWidget(self.val_thoigian_str)

    def get_ym_ctrl(self, p_idx, m_idx):
        class DummyVar: 
            def __init__(self, val): self.val = val
            def get(self): return self.val
        return {
            'type': self.DummyWidget("Trắc nghiệm"),
            'mix': self.DummyWidget("Câu hỏi & đáp án"),
            'fix': DummyVar(False),
            'start_q1': DummyVar(True),
            'pick': self.DummyWidget("9999")
        }

    def force_zero_spacing(self, p_elm, justify=False, preserve_align=False):
        pPr = p_elm.get_or_add_pPr()
        for tag in ['w:pageBreakBefore', 'w:keepNext', 'w:keepLines', 'w:contextualSpacing']:
            el = pPr.find(qn(tag))
            if el is not None:
                pPr.remove(el)
        spacing_old = pPr.find(qn('w:spacing'))
        if spacing_old is not None:
            pPr.remove(spacing_old)
        spacing_new = OxmlElement('w:spacing')
        spacing_new.set(qn('w:line'), '240')  
        spacing_new.set(qn('w:lineRule'), 'auto')
        spacing_new.set(qn('w:before'), '0')  
        spacing_new.set(qn('w:after'), '0')   
        spacing_new.set(qn('w:beforeAutospacing'), '0') 
        spacing_new.set(qn('w:afterAutospacing'), '0')  
        pPr.append(spacing_new)
        try:
            p_obj = Paragraph(p_elm, None)
            if preserve_align and p_obj.alignment in [WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.RIGHT, WD_ALIGN_PARAGRAPH.JUSTIFY]:
                pass
            else:
                p_obj.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY if justify else WD_ALIGN_PARAGRAPH.LEFT
        except: pass
        return p_elm

    def parse_youngmix_structure(self, doc_obj):
        self.youngmix_data = []
        self.youngmix_total_q = 0
        current_group = None
        tag_re = re.compile(r'^\s*[<＜](#?g\d*)[>＞]\s*', re.IGNORECASE)
        q_re = re.compile(r'^\s*(Câu|Question)\s*\d+', re.IGNORECASE)

        for block in iter_block_items(doc_obj):
            if isinstance(block, Table):
                if current_group is not None:
                    for row in block.rows:
                        for cell in row.cells:
                            for p in cell.paragraphs:
                                if q_re.match(p.text.strip().replace('\xa0', ' ').replace('\u200b', '')):
                                    current_group["q_count"] += 1
                                    self.youngmix_total_q += 1
                continue

            if not isinstance(block, Paragraph): continue
                
            p = block
            text = p.text.strip()
            if not text and not any(n.tag.endswith(('drawing', 'pict', 'object', 'OLEObject', 'oMath', 'oMathPara')) for n in p._element.iter()):
                continue

            match_tag = tag_re.match(text)
            if match_tag:
                tag_str = "<" + match_tag.group(1).lower() + ">"
                is_fixed = '#' in tag_str
                g_type = tag_str.replace('<', '').replace('>', '').replace('#', '')

                mix_type = "Câu hỏi & đáp án"
                q_type = "Trắc nghiệm"
                
                if g_type == 'g0': mix_type = "Không hoán vị"
                elif g_type == 'g1': mix_type = "Chỉ câu hỏi"
                elif g_type == 'g2': mix_type = "Chỉ đáp án"
                elif g_type == 'g3': mix_type = "Câu hỏi & đáp án"
                elif g_type == 'g4':
                    mix_type = "Chỉ câu hỏi"
                    q_type = "Tự luận"
                    is_fixed = False 
                elif g_type == 'g': mix_type = "Câu hỏi & đáp án"

                current_group = {
                    "tag": tag_str,
                    "name": f"NHÓM {len(self.youngmix_data) + 1}",
                    "q_type": q_type,
                    "mix_type": mix_type,
                    "is_fixed": is_fixed,
                    "q_count": 0,
                    "preview_data": [] 
                }
                self.youngmix_data.append(current_group)
                continue

            if current_group is None:
                current_group = {
                    "tag": "<g>",
                    "name": "NHÓM 1 (Tự động)",
                    "q_type": "Trắc nghiệm",
                    "mix_type": "Câu hỏi & đáp án",
                    "is_fixed": False,
                    "q_count": 0,
                    "preview_data": []
                }
                self.youngmix_data.append(current_group)

            if q_re.match(text):
                current_group["q_count"] += 1
                self.youngmix_total_q += 1

            is_short_answer_para = False
            if re.search(r'(?i)(đáp\s+án\s*[:\.]?)', text):
                is_short_answer_para = True
            elif re.search(r'(?:^|\n|\uFFFC)\s*A\.\s*\d', text) and not re.search(r'(?:^|\s+)B\.', text):
                is_short_answer_para = True

            has_mcq_options = bool(re.search(r'(?:^|\s+|\uFFFC+)#?[A-Da-d][\.\)](?:\s+|$)', text))
            para_chunks = [] 

            for child in p._element:
                if child.tag == qn('w:r'):
                    run = docx.text.run.Run(child, p)
                    is_ans_format = False
                    is_underlined = False
                    if run.underline and run.underline != docx.enum.text.WD_UNDERLINE.NONE:
                        is_underlined = True
                        is_ans_format = True
                        
                    rPr = child.find(qn('w:rPr'))
                    if rPr is not None:
                        color_tag = rPr.find(qn('w:color'))
                        if color_tag is not None:
                            c_val = color_tag.get(qn('w:val'))
                            if c_val and c_val.lower() != 'auto' and len(c_val) == 6:
                                try:
                                    r_color = int(c_val[0:2], 16)
                                    g_color = int(c_val[2:4], 16)
                                    b_color = int(c_val[4:6], 16)
                                    if r_color > 150 and g_color < 100 and b_color < 100:
                                        is_ans_format = True
                                except ValueError:
                                    pass

                    if run.text:
                        para_chunks.append([run.text, is_ans_format, is_underlined])
                        
                    for subchild in child:
                        tag_str = str(subchild.tag)
                        if tag_str.endswith('drawing') or tag_str.endswith('pict'):
                            para_chunks.append([" [🖼️ Hình ảnh] ", False, False])
                        elif tag_str.endswith('object') or tag_str.endswith('OLEObject'):
                            para_chunks.append([" [∑ Công thức] ", False, False])
                
                elif child.tag == qn('m:oMath') or child.tag == qn('m:oMathPara'):
                    math_texts = []
                    for t in child.iter(qn('m:t')):
                        if t.text: math_texts.append(t.text)
                    if math_texts:
                        para_chunks.append(["".join(math_texts), False, False])
                    else:
                        para_chunks.append([" [∑ Công thức] ", False, False])

            if is_short_answer_para:
                for chunk in para_chunks: chunk[1] = True
            elif has_mcq_options:
                full_para_text = "".join(chunk[0] for chunk in para_chunks)
                pattern = r'(?:^|\s+|\uFFFC+)#?([A-Da-d][\.\)])'
                matches = list(re.finditer(pattern, full_para_text))
                
                if matches:
                    char_is_ans = []
                    char_is_underlined = []
                    for chunk in para_chunks:
                        char_is_ans.extend([chunk[1]] * len(chunk[0]))
                        char_is_underlined.extend([chunk[2]] * len(chunk[0]))
                        
                    boundaries = [m.start(1) for m in matches]
                    spans = [0]
                    for b in boundaries: spans.append(b)
                    spans.append(len(full_para_text))
                    
                    marked_regions = set()
                    for i, m in enumerate(matches):
                        label_start = m.start(1)
                        label_end = m.end(1)
                        is_label_marked = any(char_is_ans[idx] for idx in range(label_start, label_end) if idx < len(char_is_ans))
                        if is_label_marked:
                            marked_regions.add(i + 1)
                            
                    new_para_chunks = []
                    for i in range(len(full_para_text)):
                        char = full_para_text[i]
                        region_idx = 0
                        for r in range(len(spans) - 1):
                            if i >= spans[r] and i < spans[r+1]:
                                region_idx = r
                                break
                        is_red = (region_idx in marked_regions)
                        is_und = char_is_underlined[i]
                        
                        if not new_para_chunks:
                            new_para_chunks.append([char, is_red, is_und])
                        else:
                            if is_red == new_para_chunks[-1][1] and is_und == new_para_chunks[-1][2]:
                                new_para_chunks[-1][0] += char
                            else:
                                new_para_chunks.append([char, is_red, is_und])
                                
                    para_chunks = new_para_chunks
            else:
                for chunk in para_chunks: chunk[1] = False

            for chunk in para_chunks:
                current_group["preview_data"].append((chunk[0], chunk[1], chunk[2]))

            current_group["preview_data"].append(("\n", False, False))
            
        if len(self.youngmix_data) > 0:
            first_group = self.youngmix_data[0]
            if first_group["name"] == "NHÓM 1 (Tự động)" and first_group["q_count"] == 0:
                self.youngmix_data.pop(0) 
                for i, g in enumerate(self.youngmix_data):
                    if g["name"].startswith("NHÓM"): 
                        g["name"] = f"NHÓM {i + 1}"

    def convert_auto_numbering(self, doc_obj):
        try:
            numbering_part = None
            try:
                numbering_part = doc_obj.part.numbering_part
            except: pass
            if not numbering_part: return
            
            num_to_abs = {}
            for num in numbering_part._element.xpath('.//w:num'):
                numId = num.get(qn('w:numId'))
                abs_el = num.find(qn('w:abstractNumId'))
                if abs_el is not None:
                    num_to_abs[numId] = abs_el.get(qn('w:val'))
                    
            abs_to_lvl = {}
            for absNum in numbering_part._element.xpath('.//w:abstractNum'):
                absId = absNum.get(qn('w:abstractNumId'))
                for lvl in absNum.xpath('.//w:lvl'):
                    ilvl = lvl.get(qn('w:ilvl')) or '0'
                    numFmt_el = lvl.find(qn('w:numFmt'))
                    numFmt = numFmt_el.get(qn('w:val')) if numFmt_el is not None else 'decimal'
                    lvlText_el = lvl.find(qn('w:lvlText'))
                    lvlText = lvlText_el.get(qn('w:val')) if lvlText_el is not None else ''
                    abs_to_lvl[(absId, ilvl)] = {'fmt': numFmt, 'text': lvlText}
            
            counters = {}
            
            def format_num(val, fmt):
                if fmt == 'decimal': return str(val)
                elif fmt in ('upperLetter', 'lowerLetter'): 
                    c = chr(64 + val) if fmt == 'upperLetter' else chr(96 + val)
                    return c if 1 <= val <= 26 else str(val)
                elif fmt in ('upperRoman', 'lowerRoman'):
                    roman_map = ((1000, 'M'), (900, 'CM'), (500, 'D'), (400, 'CD'), (100, 'C'), (90, 'XC'), (50, 'L'), (40, 'XL'), (10, 'X'), (9, 'IX'), (5, 'V'), (4, 'IV'), (1, 'I'))
                    res = ""
                    n = val
                    for dec, rom in roman_map:
                        while n >= dec:
                            res += rom
                            n -= dec
                    return res.lower() if fmt == 'lowerRoman' else res
                else: return str(val)
            
            for p_elm in doc_obj._body._body.xpath('.//w:p'):
                texts = []
                for t in p_elm.iter(qn('w:t')):
                    if t.text: texts.append(t.text)
                p_text = "".join(texts).strip()
                if re.search(r'(?:Câu|Question)\s*\d+', p_text, re.IGNORECASE):
                    counters.clear()
                
                has_num = False
                pPr = p_elm.find(qn('w:pPr'))
                if pPr is not None:
                    numPr = pPr.find(qn('w:numPr'))
                    if numPr is not None:
                        numId_el = numPr.find(qn('w:numId'))
                        ilvl_el = numPr.find(qn('w:ilvl'))
                        if numId_el is not None and numId_el.get(qn('w:val')) != '0':
                            has_num = True
                            numId = numId_el.get(qn('w:val'))
                            ilvl = ilvl_el.get(qn('w:val')) if ilvl_el is not None else '0'
                            
                            absId = num_to_abs.get(numId)
                            if absId:
                                lvl_info = abs_to_lvl.get((absId, ilvl))
                                if lvl_info:
                                    key = (numId, ilvl)
                                    if key not in counters: counters[key] = 0
                                    counters[key] += 1
                                    
                                    for k in list(counters.keys()):
                                        if k[0] == numId and int(k[1]) > int(ilvl):
                                            counters[k] = 0
                                            
                                    val = counters[key]
                                    fmt = lvl_info['fmt']
                                    txt_template = lvl_info['text']
                                    
                                    num_str = format_num(val, fmt)
                                    level_idx_str = f"%{int(ilvl)+1}"
                                    rendered_text = txt_template.replace(level_idx_str, num_str)
                                    rendered_text = re.sub(r'%\d+', '', rendered_text)
                                    
                                    pPr.remove(numPr)
                                    pStyle = pPr.find(qn('w:pStyle'))
                                    if pStyle is not None:
                                        pStyle.set(qn('w:val'), 'Normal')
                                    
                                    r_new = OxmlElement('w:r')
                                    rPr_new = OxmlElement('w:rPr')
                                    b = OxmlElement('w:b')
                                    b.set(qn('w:val'), '1')
                                    rPr_new.append(b)
                                    r_new.append(rPr_new)
                                    
                                    t = OxmlElement('w:t')
                                    t.set(qn('xml:space'), 'preserve')
                                    t.text = rendered_text + " "
                                    r_new.append(t)
                                    
                                    first_r = p_elm.find(qn('w:r'))
                                    if first_r is not None:
                                        first_r.addprevious(r_new)
                                    else:
                                        p_elm.append(r_new)
                
                if not has_num and p_text:
                    counters.clear()
        except Exception: pass

    def split_soft_returns(self, doc_obj):
        body = doc_obj._body._body
        ps = []
        for p in body.xpath('.//w:p'): ps.append(p)
        for p in ps:
            if not p.xpath('.//w:br'): continue
            
            parent = p.getparent()
            if parent is None: continue
            
            new_paras = []
            current_p = OxmlElement('w:p')
            pPr = p.find(qn('w:pPr'))
            if pPr is not None: current_p.append(copy.deepcopy(pPr))
            
            for child in p:
                if child.tag == qn('w:pPr'): continue
                if child.tag == qn('w:r'):
                    current_r = OxmlElement('w:r')
                    rPr = child.find(qn('w:rPr'))
                    if rPr is not None: current_r.append(copy.deepcopy(rPr))
                    
                    r_has_content = False
                    for subchild in child:
                        if subchild.tag == qn('w:rPr'): continue
                        elif subchild.tag == qn('w:br'):
                            if r_has_content: current_p.append(current_r)
                            new_paras.append(current_p)
                            
                            current_p = OxmlElement('w:p')
                            if pPr is not None: current_p.append(copy.deepcopy(pPr))
                            current_r = OxmlElement('w:r')
                            if rPr is not None: current_r.append(copy.deepcopy(rPr))
                            r_has_content = False
                        else:
                            current_r.append(copy.deepcopy(subchild))
                            r_has_content = True
                    if r_has_content: current_p.append(current_r)
                else:
                    current_p.append(copy.deepcopy(child))
            
            new_paras.append(current_p)
            for np in new_paras: parent.insert(parent.index(p), np)
            parent.remove(p)

    def unwrap_omathpara(self, p_elm):
        if p_elm.tag != qn('w:p'): return p_elm
        for math_para in list(p_elm.findall(qn('m:oMathPara'))):
            parent = math_para.getparent()
            if parent is not None:
                for math_node in list(math_para):
                    math_para.addprevious(math_node)
                parent.remove(math_para)
        return p_elm

    def reset_paragraph_indent(self, p_elm, left_cm=0, first_line_cm=0):
        pPr = p_elm.get_or_add_pPr()
        ind = pPr.find(qn('w:ind'))
        if ind is None:
            ind = OxmlElement('w:ind')
            pPr.append(ind)
            
        for attr in ('left', 'right', 'firstLine', 'hanging', 'leftChars', 'rightChars', 'firstLineChars', 'hangingChars'):
            if qn(f'w:{attr}') in ind.attrib:
                del ind.attrib[qn(f'w:{attr}')]
                
        ind.set(qn('w:left'), str(int(left_cm * 567)))
        ind.set(qn('w:firstLine'), str(int(first_line_cm * 567)))
        ind.set(qn('w:hanging'), '0')

    def force_part2_indent(self, p_elm, doc_obj):
        pPr = p_elm.get_or_add_pPr()
        for tag in ('w:ind', 'w:numPr'):
            for el in pPr.findall(qn(tag)):
                pPr.remove(el)
                
        new_ind = OxmlElement('w:ind')
        new_ind.set(qn('w:left'), '0')
        new_ind.set(qn('w:firstLine'), '283')
        new_ind.set(qn('w:hanging'), '0')
        pPr.append(new_ind)
        
        para = Paragraph(p_elm, doc_obj)
        try:
            para.style = doc_obj.styles['Normal']
        except: pass
        para.paragraph_format.left_indent = Cm(0)
        para.paragraph_format.first_line_indent = Cm(0.5)

    def clean_empty_lines(self, p_elm):
        if p_elm.tag != qn('w:p'): return p_elm
        for child in list(p_elm):
            if child.tag == qn('w:r'):
                stop = False
                for sub in list(child):
                    tag = str(sub.tag)
                    if tag.endswith('}br') or tag.endswith('}tab'):
                        child.remove(sub)
                    elif tag.endswith('}t'):
                        if sub.text:
                            sub.text = sub.text.lstrip()
                            if sub.text: stop = True; break
                    elif any(tag.endswith(t) for t in ('}drawing', '}object', '}pict', '}AlternateContent')):
                        stop = True; break
                if stop: break
            elif child.tag in (qn('m:oMath'), qn('m:oMathPara')): break 

        for child in reversed(list(p_elm)):
            if child.tag == qn('w:r'):
                stop = False
                for sub in reversed(list(child)):
                    tag = str(sub.tag)
                    if tag.endswith('}br') or tag.endswith('}tab'):
                        child.remove(sub)
                    elif tag.endswith('}t'):
                        if sub.text:
                            sub.text = sub.text.rstrip()
                            if sub.text: stop = True; break
                    elif any(tag.endswith(t) for t in ('}drawing', '}object', '}pict', '}AlternateContent')):
                        stop = True; break
                if stop: break
            elif child.tag in (qn('m:oMath'), qn('m:oMathPara')): break 

        br_count = 0
        for child in list(p_elm):
            if child.tag == qn('w:r'):
                for sub in list(child):
                    tag = str(sub.tag)
                    if tag.endswith('}br'):
                        br_count += 1
                        if br_count > 1: child.remove(sub)
                    elif tag.endswith('}t'):
                        if sub.text and sub.text.strip(): br_count = 0
                    elif any(tag.endswith(t) for t in ('}drawing', '}object', '}pict', '}AlternateContent')):
                        br_count = 0
            elif child.tag in (qn('m:oMath'), qn('m:oMathPara')): br_count = 0
        return p_elm

    def _build_mapping(self, p_elm):
        mapping = []
        full_text = ""
        for child in p_elm:
            if child.tag == qn('w:pPr'): continue
            elif child.tag == qn('w:r'):
                for subchild in child:
                    if subchild.tag == qn('w:rPr'): continue
                    elif subchild.tag == qn('w:t'):
                        if subchild.text:
                            for i in range(len(subchild.text)): mapping.append({'node': subchild, 'type': 'text', 'idx': i})
                            full_text += subchild.text
                    elif subchild.tag == qn('w:tab'):
                        mapping.append({'node': subchild, 'type': 'tab'}); full_text += "\t"
                    elif subchild.tag == qn('w:br'):
                        mapping.append({'node': subchild, 'type': 'br'}); full_text += "\n"
                    else:
                        mapping.append({'node': subchild, 'type': 'element'}); full_text += "\uFFFC"
            elif child.tag in (qn('m:oMath'), qn('m:oMathPara')):
                m_texts = []
                for t in child.iter(qn('m:t')):
                    if t.text: m_texts.append(t.text)
                math_text = "".join(m_texts)
                if math_text:
                    for i in range(len(math_text)):
                        mapping.append({'node': child, 'type': 'element_text', 'idx': i})
                    full_text += math_text
                else:
                    mapping.append({'node': child, 'type': 'element'})
                    full_text += "\uFFFC"
            elif child.tag in (qn('w:hyperlink'), qn('w:smartTag')):
                for sub_r in child.findall(qn('w:r')):
                    for subchild in sub_r:
                        if subchild.tag == qn('w:rPr'): continue
                        elif subchild.tag == qn('w:t'):
                            if subchild.text:
                                for i in range(len(subchild.text)): mapping.append({'node': subchild, 'type': 'text', 'idx': i})
                                full_text += subchild.text
                        elif subchild.tag == qn('w:tab'):
                            mapping.append({'node': subchild, 'type': 'tab'}); full_text += "\t"
                        elif subchild.tag == qn('w:br'):
                            mapping.append({'node': subchild, 'type': 'br'}); full_text += "\n"
                        else:
                            mapping.append({'node': subchild, 'type': 'element'}); full_text += "\uFFFC"
            else:
                mapping.append({'node': child, 'type': 'element'}); full_text += "\uFFFC"
        return mapping, full_text

    def extract_options_from_paragraph(self, p_elm, current_ans_count, is_part2=False):
        expected_labels = ('a', 'b', 'c', 'd') if is_part2 else ('A', 'B', 'C', 'D')
        opts_with_status = []
        c_mapping, full_text = self._build_mapping(p_elm)
                
        if not full_text.strip():
            temp_nodes = []
            for c in p_elm:
                if str(c.tag) != qn('w:pPr'): temp_nodes.append(copy.deepcopy(c))
            return [(temp_nodes, False, False)]

        split_points = []
        found_count = current_ans_count
        search_start = 0
        
        while found_count < 4:
            next_label = expected_labels[found_count]
            if is_part2:
                pat = r'(^\s*|\t+|\n+|\uFFFC+|\s{2,}|(?:Câu|Question)\s*\d+[\.\:]\s+)(#?)(' + next_label + r'[\.\)])'
                match = re.search(pat, full_text[search_start:], re.IGNORECASE)
            else:
                pat = r'(^\s*|\t+|\n+|\uFFFC+|\s{2,}|(?:Câu|Question)\s*\d+[\.\:]\s+)(#?)(' + next_label + r'[\.\)])'
                match = re.search(pat, full_text[search_start:]) 
            
            if match:
                is_fixed = (match.group(2) == '#')
                start_idx_for_check = search_start + match.start(2) if match.group(2) else search_start + match.start(3)
                end_idx_for_check = search_start + match.end()
                
                is_true = False
                for idx in range(start_idx_for_check, end_idx_for_check):
                    if idx < len(c_mapping):
                        target_node = c_mapping[idx]['node']
                        parent_r = target_node.getparent()
                        if parent_r is not None and str(parent_r.tag).endswith('}r'):
                            rPr = parent_r.find(qn('w:rPr'))
                            if rPr is not None:
                                u_tag = rPr.find(qn('w:u'))
                                if u_tag is not None and u_tag.get(qn('w:val')) != 'none': 
                                    is_true = True
                                    break
                                color_tag = rPr.find(qn('w:color'))
                                if color_tag is not None:
                                    c_val = color_tag.get(qn('w:val'))
                                    if c_val and c_val.lower() != 'auto' and len(c_val) == 6:
                                        try:
                                            r_color = int(c_val[0:2], 16)
                                            g_color = int(c_val[2:4], 16)
                                            b_color = int(c_val[4:6], 16)
                                            if r_color > 150 and g_color < 100 and b_color < 100:
                                                is_true = True
                                                break
                                        except ValueError:
                                            pass
                
                question_match = re.search(r'(^|[\s\t\n\uFFFC]+)(?:Câu|Question)\s*\d+(\([^)]+\))?[\.\:\s]', full_text[search_start:], re.IGNORECASE)
                if question_match and question_match.start() < match.start():
                    if search_start == 0 and question_match.start() == 0:
                        text_before_match = full_text[search_start:search_start + match.start()]
                        text_after_first_q = text_before_match[question_match.end():]
                        if re.search(r'(^|[\s\t\n\uFFFC]+)(?:Câu|Question)\s*\d+(\([^)]+\))?[\.\:\s]', text_after_first_q, re.IGNORECASE):
                            break
                    else:
                        break

                split_start_idx = search_start + match.start(2) if match.group(2) else search_start + match.start(3)
                split_points.append((split_start_idx, is_true, is_fixed))
                search_start = search_start + match.end()
                found_count += 1
            else: break
                
        if not split_points:
            temp_nodes = []
            for c in p_elm:
                if str(c.tag) != qn('w:pPr'): temp_nodes.append(copy.deepcopy(c))
            return [(temp_nodes, False, False)]
            
        ranges = []
        for i in range(len(split_points)):
            start = split_points[i][0]
            is_true = split_points[i][1]
            is_fixed = split_points[i][2]
            end = split_points[i+1][0] if i+1 < len(split_points) else len(full_text)
            ranges.append((start, end, is_true, is_fixed))
            
        for start_idx, end_idx, is_true, is_fixed in ranges:
            p_clone = copy.deepcopy(p_elm)
            clone_mapping, _ = self._build_mapping(p_clone)
            
            nodes_to_remove = set()
            node_texts = {}
            for i, item in enumerate(clone_mapping):
                if i < start_idx or i >= end_idx:
                    if item['type'] in ('tab', 'br', 'element', 'element_text'):
                        nodes_to_remove.add(item['node'])
                    elif item['type'] == 'text':
                        if item['node'] not in node_texts:
                            node_texts[item['node']] = list(item['node'].text)
                        node_texts[item['node']][item['idx']] = ""
                        
            for node, chars in node_texts.items():
                node.text = "".join(chars)
            for node in nodes_to_remove:
                if node.getparent() is not None:
                    node.getparent().remove(node)
            
            opt_elements = []
            for c in p_clone:
                if str(c.tag) != qn('w:pPr'): opt_elements.append(c)
                
            opts_with_status.append((opt_elements, is_true, is_fixed))
            
        return opts_with_status

    def extract_and_clean_dapan(self, elm):
        ans_val = None
        p_elms = []
        for p in elm.iter(qn('w:p')): p_elms.append(p)
            
        for p_elm in p_elms:
            c_mapping, full_text = self._build_mapping(p_elm)
            match = re.search(r'(?i)(đáp\s+án\s*[:\.]|(?:^|\n|\uFFFC)\s*A\.)\s*(.*)', full_text, re.IGNORECASE)
            
            if match and re.search(r'\d', match.group(2)):
                if not ans_val: ans_val = match.group(2).strip()
                start_idx = match.start(1)
                
                nodes_to_remove = set()
                node_texts = {}
                
                for i in range(start_idx, len(full_text)):
                    item = c_mapping[i]
                    if item['type'] in ('tab', 'br', 'element', 'element_text'):
                        nodes_to_remove.add(item['node'])
                    elif item['type'] == 'text':
                        if item['node'] not in node_texts:
                            node_texts[item['node']] = list(item['node'].text)
                        node_texts[item['node']][item['idx']] = ""
                        
                for node, chars in node_texts.items(): node.text = "".join(chars)
                for node in nodes_to_remove:
                    if node.getparent() is not None: node.getparent().remove(node)
                    
            texts = []
            for t in p_elm.iter(qn('w:t')):
                if t.text: texts.append(t.text)
            new_text = "".join(texts).strip()
            
            has_obj = False
            for n in p_elm.iter():
                if str(n.tag).endswith('imagedata') or str(n.tag).endswith('pic') or str(n.tag).endswith('object') or str(n.tag).endswith('drawing') or str(n.tag).endswith('oMath'):
                    has_obj = True
                    break
            
            if not new_text and not has_obj:
                if p_elm.getparent() is not None: p_elm.getparent().remove(p_elm)
        return ans_val

    def parse_document_structure(self, doc_obj):
        parsed = {"header": [], "parts": []}; errors = []
        part_re = re.compile(r'^[\s\W]*PHẦN\s+([IVX]+)', re.IGNORECASE)
        muc_re = re.compile(r'^[\s\W]*([IVX]+)\s*\.\s*(.*)', re.IGNORECASE)
        q_re = re.compile(r'^[\s\xa0\u200b]*(#?)\s*(?:Câu|Question)\s*\d+\s*(\([^)]+\))?[\.\:\s]', re.IGNORECASE)
        current_part = current_muc = current_question = None; is_header = True
        has_explicit_part = False
        is_ym_mode = getattr(self, 'youngmix_var', None) and self.youngmix_var.get()
        
        blocks = []
        for b in iter_block_items(doc_obj): blocks.append(b)
        
        for block_idx, block in enumerate(blocks):
            is_table = isinstance(block, Table)
            txt = ""
            if not is_table:
                txt = block.text.strip().replace('\xa0', ' ').replace('\u200b', '')
            else:
                for row in block.rows:
                    for cell in row.cells:
                        for p in cell.paragraphs:
                            p_text = p.text.strip().replace('\xa0', ' ').replace('\u200b', '')
                            if p_text:
                                txt = p_text
                                break
                        if txt: break
                    if txt: break
            
            has_obj = False
            if not is_table:
                for n in block._element.iter():
                    tag_str = str(n.tag)
                    if tag_str.endswith('}imagedata') or tag_str.endswith('}pic') or tag_str.endswith('}object') or tag_str.endswith('}drawing') or tag_str.endswith('}oMath') or tag_str.endswith('}AlternateContent'):
                        has_obj = True
                        break
                        
            if not txt and not has_obj and not is_table: continue
            
            m_part_original = part_re.match(txt)
            m_part = m_part_original
            m_muc = muc_re.match(txt)
            m_q = q_re.match(txt)
            
            if is_ym_mode:
                m_part = None 
                m_ym = re.match(r'^\s*[<＜](#?g\d*)[>＞]\s*', txt, re.IGNORECASE)
                if m_ym:
                    is_muc_match = True
                    m_muc = None 
                else:
                    is_muc_match = False
            else:
                m_ym = None
                is_muc_match = False
                if m_muc and "Câu" not in txt and "Question" not in txt and not txt.upper().startswith("PHẦN"):
                    is_muc_match = True
                
            is_q_match = False
            if m_q:
                is_q_match = True
                
            if m_part_original or is_muc_match or is_q_match:
                is_header = False
                
            if is_header:
                parsed["header"].append(copy.deepcopy(block._element))
                continue
            
            if m_part:
                has_explicit_part = True
                t_up = txt.upper()
                p_type = 1
                if "PHẦN IV" in t_up: p_type = 4
                elif "PHẦN III" in t_up: p_type = 3
                elif "PHẦN II" in t_up: p_type = 2
                
                current_part = {"title": txt, "title_element": copy.deepcopy(block._element), "type": p_type, "mucs": [], "is_explicit": True}
                parsed["parts"].append(current_part)
                current_muc = None
                current_question = None
                continue
                
            if is_muc_match:
                if current_part is None:
                    current_part = {"title": "PHẦN I (Tự nhận diện)", "title_element": None, "type": 1, "mucs": [], "is_explicit": False}
                    parsed["parts"].append(current_part)
                    
                if is_ym_mode and m_ym:
                    muc_title = "<" + m_ym.group(1).lower() + ">"
                    is_fixed_muc = '#' in muc_title
                    muc_title_element = copy.deepcopy(block._element)
                else:
                    muc_title = m_muc.group(2).strip()
                    is_fixed_muc = txt.lstrip().startswith('#')
                    muc_title_element = copy.deepcopy(block._element)
                    if is_fixed_muc:
                        for t in muc_title_element.iter(qn('w:t')):
                            if t.text and '#' in t.text:
                                t.text = t.text.replace('#', '', 1)
                                break

                current_muc = {"title": muc_title, "title_element": muc_title_element, "description_elements": [], "questions": [], "is_fixed": is_fixed_muc}
                current_part["mucs"].append(current_muc)
                current_question = None
                continue
                
            if not current_muc and current_part is not None:
                current_muc = {"title": None, "title_element": None, "description_elements": [], "questions": [], "is_fixed": False}
                current_part["mucs"].append(current_muc)
                
            if is_q_match:
                if current_part is None:
                    current_part = {"title": "PHẦN I (Tự nhận diện)", "title_element": None, "type": 1, "mucs": [], "is_explicit": False}
                    parsed["parts"].append(current_part)
                
                if not current_muc:
                    current_muc = {"title": None, "title_element": None, "description_elements": [], "questions": [], "is_fixed": False}
                    current_part["mucs"].append(current_muc)
                    
                is_fixed_q = False
                if m_q.group(1): is_fixed_q = True
                
                q_elem_copy = copy.deepcopy(block._element)
                if is_fixed_q:
                    for t in q_elem_copy.iter(qn('w:t')):
                        if t.text and '#' in t.text:
                            t.text = t.text.replace('#', '', 1)
                            break
                            
                current_question = {'raw_text': txt, 'q_elements': [q_elem_copy], 'answers': [], 'is_fixed': is_fixed_q}
                current_muc["questions"].append(current_question)

                if is_ym_mode:
                    has_ABCD = False
                    has_abcd = False
                    for f_idx in range(block_idx, len(blocks)):
                        f_b = blocks[f_idx]
                        f_is_tbl = isinstance(f_b, Table)
                        f_txt = f_b.text.strip().replace('\xa0', ' ').replace('\u200b', '') if not f_is_tbl else ""
                        
                        if f_idx > block_idx:
                            if q_re.match(f_txt) or part_re.match(f_txt) or re.match(r'^\s*[<＜](#?g\d*)[>＞]\s*', f_txt, re.IGNORECASE):
                                break
                        
                        if re.search(r'(^\s*|\t+|\n+|\uFFFC+|\s{2,})#?[A-D][\.\)]', f_txt):
                            has_ABCD = True
                        if re.search(r'(^\s*|\t+|\n+|\uFFFC+|\s{2,})#?[a-d][\.\)]', f_txt):
                            has_abcd = True
                            
                    if has_ABCD:
                        current_question['detected_part2'] = False
                        current_question['format_locked'] = True
                    elif has_abcd:
                        current_question['detected_part2'] = True
                        current_question['format_locked'] = True
                
                should_process_options = is_ym_mode or (current_part and current_part["type"] in (1, 2))
                if should_process_options:
                    is_part2 = False
                    if not is_ym_mode and current_part and current_part["type"] == 2:
                        is_part2 = True
                        
                    if is_ym_mode:
                        is_part2 = current_question.get('detected_part2', False)
                    
                    expected_labels = ['a', 'b', 'c', 'd'] if is_part2 else ['A', 'B', 'C', 'D']
                    
                    match = None
                    if is_part2:
                        pat = r'(^\s*|\t+|\n+|\uFFFC+|\s{2,}|(?:Câu|Question)\s*\d+[\.\:]\s+)#?(' + expected_labels[0] + r'[\.\)])'
                        match = re.search(pat, txt, re.IGNORECASE)
                    else:
                        pat = r'(^\s*|\t+|\n+|\uFFFC+|\s{2,}|(?:Câu|Question)\s*\d+[\.\:]\s+)#?(' + expected_labels[0] + r'[\.\)])'
                        match = re.search(pat, txt)
                        
                    if match:
                        opts_with_status = self.extract_options_from_paragraph(block._element, 0, is_part2)
                        if opts_with_status and len(opts_with_status[0][0]) > 0:
                            q_elm = current_question['q_elements'][0]
                            c_mapping, full_text_q = self._build_mapping(q_elm)
                            
                            q_match = None
                            if is_part2:
                                q_pat = r'(^\s*|\t+|\n+|\uFFFC+|\s{2,}|(?:Câu|Question)\s*\d+[\.\:]\s+)(#?)(' + expected_labels[0] + r'[\.\)])'
                                q_match = re.search(q_pat, full_text_q, re.IGNORECASE)
                            else:
                                q_pat = r'(^\s*|\t+|\n+|\uFFFC+|\s{2,}|(?:Câu|Question)\s*\d+[\.\:]\s+)(#?)(' + expected_labels[0] + r'[\.\)])'
                                q_match = re.search(q_pat, full_text_q)
                                
                            if q_match:
                                first_label_start = q_match.start(2)
                                if not q_match.group(2): first_label_start = q_match.start(3)
                                
                                nodes_to_remove = set()
                                node_texts = {}
                                for i, item in enumerate(c_mapping):
                                    if i >= first_label_start: 
                                        itype = item['type']
                                        inode = item['node']
                                        if itype in ('tab', 'br', 'element', 'element_text'):
                                            nodes_to_remove.add(inode)
                                        elif itype == 'text':
                                            if inode not in node_texts:
                                                node_texts[inode] = list(inode.text)
                                            node_texts[inode][item['idx']] = ""
                                            
                                for node, chars in node_texts.items():
                                    node.text = "".join(chars)
                                for node in nodes_to_remove:
                                    if node.getparent() is not None:
                                        node.getparent().remove(node)
                            
                            for opt_elements, is_true, is_fixed in opts_with_status:
                                if len(current_question['answers']) >= 4:
                                    break 
                                new_p = OxmlElement('w:p')
                                for elm in opt_elements:
                                    new_p.append(elm)
                                current_question['answers'].append({"elements": [new_p], "is_true": is_true, "is_fixed": is_fixed})
                continue
            
            if current_question:
                if is_table:
                    if is_ym_mode:
                        for row in block.rows:
                            for cell in row.cells:
                                for p in cell.paragraphs:
                                    if q_re.match(p.text.strip().replace('\xa0', ' ').replace('\u200b', '')):
                                        dummy_q = {'raw_text': p.text.strip(), 'q_elements': [], 'answers': [], 'is_fixed': True, 'ym_type': 4, 'is_virtual': True}
                                        current_muc["questions"].append(dummy_q)
                    if len(current_question['answers']) > 0:
                        current_question['answers'][-1]["elements"].append(copy.deepcopy(block._element))
                    else:
                        current_question['q_elements'].append(copy.deepcopy(block._element))
                    continue
                    
                should_process_options = is_ym_mode or (current_part and current_part["type"] in (1, 2))
                if should_process_options:
                    expected_idx = len(current_question['answers'])
                    is_part2 = False
                    if not is_ym_mode and current_part and current_part["type"] == 2:
                        is_part2 = True
                        
                    if is_ym_mode:
                        is_part2 = current_question.get('detected_part2', False)
                    
                    is_ans = False
                    if expected_idx < 4:
                        labels = ['a', 'b', 'c', 'd'] if is_part2 else ['A', 'B', 'C', 'D']
                        next_label = labels[expected_idx]
                        
                        if is_part2:
                            pat = r'(^\s*|\t+|\n+|\uFFFC+|\s{2,}|(?:Câu|Question)\s*\d+[\.\:]\s+)#?(' + next_label + r'[\.\)])'
                            is_ans = bool(re.search(pat, txt, re.IGNORECASE))
                        else:
                            pat = r'(^\s*|\t+|\n+|\uFFFC+|\s{2,}|(?:Câu|Question)\s*\d+[\.\:]\s+)#?(' + next_label + r'[\.\)])'
                            is_ans = bool(re.search(pat, txt))
                            
                    if is_ans:
                        p_copy = copy.deepcopy(block._element)
                        num_prs = []
                        for n in p_copy.iter(qn('w:numPr')): num_prs.append(n)
                        for n in num_prs:
                            if n.getparent() is not None:
                                n.getparent().remove(n)
                        opts_with_status = self.extract_options_from_paragraph(p_copy, expected_idx, is_part2)
                        for opt_elements, is_true, is_fixed in opts_with_status:
                            if len(current_question['answers']) >= 4:
                                break 
                            new_p = OxmlElement('w:p')
                            for elm in opt_elements:
                                new_p.append(elm)
                            current_question['answers'].append({"elements": [new_p], "is_true": is_true, "is_fixed": is_fixed})
                    elif len(current_question['answers']) > 0: 
                        current_question['answers'][-1]["elements"].append(copy.deepcopy(block._element))
                    else: 
                        current_question['q_elements'].append(copy.deepcopy(block._element))
                else:
                    current_question['q_elements'].append(copy.deepcopy(block._element))
            elif current_muc is not None:
                if "description_elements" not in current_muc:
                    current_muc["description_elements"] = []
                if is_table:
                    if is_ym_mode:
                        for row in block.rows:
                            for cell in row.cells:
                                for p in cell.paragraphs:
                                    if q_re.match(p.text.strip().replace('\xa0', ' ').replace('\u200b', '')):
                                        dummy_q = {'raw_text': p.text.strip(), 'q_elements': [], 'answers': [], 'is_fixed': True, 'ym_type': 4, 'is_virtual': True}
                                        current_muc["questions"].append(dummy_q)
                current_muc["description_elements"].append(copy.deepcopy(block._element))
        
        if not has_explicit_part and len(parsed["parts"]) > 0:
            parsed["parts"][0]["is_explicit"] = False
            
        for part in parsed["parts"]:
            for muc in part["mucs"]:
                for q in muc["questions"]:
                    num_ans = len(q["answers"])
                    has_dapan_text = False
                    
                    for elm in q["q_elements"]:
                        if elm.tag == qn('w:p'):
                            _, f_text = self._build_mapping(elm)
                            if re.search(r'(?i)(đáp\s+án\s*[:\.]?|(?:^|\n|\uFFFC)\s*[A-D]\.)\s*(.*)', f_text):
                                has_dapan_text = True
                                break
                    
                    if is_ym_mode:
                        if num_ans >= 2:
                            q['ym_type'] = 2 if q.get('detected_part2') else 1
                        elif num_ans == 1:
                            for a in q['answers']:
                                q['q_elements'].extend(a['elements'])
                            q['answers'] = [] 
                            q['ym_type'] = 3
                        else:
                            if has_dapan_text:
                                q['ym_type'] = 3
                            else:
                                q['ym_type'] = 4
                    else:
                        if num_ans == 4:
                            q['ym_type'] = part.get("type", 1) 
                        elif num_ans > 0:
                            for a in q['answers']:
                                q['q_elements'].extend(a['elements'])
                            q['answers'] = [] 
                            
                            has_dapan_now = False
                            for elm in q["q_elements"]:
                                if elm.tag == qn('w:p'):
                                    _, f_text = self._build_mapping(elm)
                                    if re.search(r'(?i)(đáp\s+án\s*[:\.]?|(?:^|\n|\uFFFC)\s*[A-D]\.)\s*(.*)', f_text):
                                        has_dapan_now = True
                                        break
                                        
                            if has_dapan_now:
                                q['ym_type'] = 3
                            else:
                                q['ym_type'] = 4 
                        else:
                            if has_dapan_text:
                                q['ym_type'] = 3 
                            else:
                                q['ym_type'] = 4 

        if not is_ym_mode:
            for part in parsed["parts"]:
                if part.get("type") == 3:
                    has_answer = False
                    for muc in part.get("mucs", []):
                        for q in muc.get("questions", []):
                            if q.get('ym_type') == 3:
                                has_answer = True
                                break
                        if has_answer: break
                    if not has_answer:
                        part["type"] = 4 

        return parsed, errors

    def preprocess_textboxes(self, doc_obj):
        body = doc_obj._body._body
        changed = True
        loop_counter = 0
        while changed and loop_counter < 100:
            loop_counter += 1
            changed = False
            
            boxes = []
            for txbx in body.xpath('.//*[local-name()="txbxContent"] | .//*[local-name()="textbox"]'):
                boxes.append(txbx)
                
            for txbx in boxes:
                parent = txbx.getparent()
                while parent is not None and parent.tag not in (qn('w:p'), qn('w:tbl')):
                    parent = parent.getparent()
                
                if parent is not None:
                    content_node = txbx
                    if txbx.tag != qn('w:txbxContent'):
                        content_node = txbx.find(qn('w:txbxContent'))
                        
                    if content_node is not None:
                        extracted = False
                        children = []
                        for child in content_node: children.append(child)
                        for child in children:
                            if child.tag in (qn('w:p'), qn('w:tbl')):
                                parent.addprevious(copy.deepcopy(child))
                                extracted = True
                        if extracted:
                            for c in children: content_node.remove(c)
                            wrapper = txbx
                            while wrapper is not None and not (str(wrapper.tag).endswith('drawing') or str(wrapper.tag).endswith('shape')):
                                wrapper = wrapper.getparent()
                            if wrapper is not None and wrapper.getparent() is not None:
                                wrapper.getparent().remove(wrapper)
                            changed = True
                            break

    def preprocess_layout_tables(self, doc_obj):
        body = doc_obj._body._body
        changed = True
        loop_counter = 0
        is_ym_mode = getattr(self, 'youngmix_var', None) and self.youngmix_var.get()
        
        while changed and loop_counter < 100:
            loop_counter += 1
            changed = False
            tbls = []
            for tbl in body.xpath('.//*[local-name()="tbl"]'): tbls.append(tbl)
            
            for tbl in tbls:
                if tbl.getparent() is None: continue
                table = Table(tbl, doc_obj)
                
                has_media = False
                for row in table.rows:
                    for cell in row.cells:
                        for p in cell.paragraphs:
                            for n in p._element.iter():
                                tag = str(n.tag)
                                if tag.endswith('}imagedata') or tag.endswith('}pic') or tag.endswith('}drawing') or tag.endswith('}object'):
                                    has_media = True
                                    break
                            if has_media: break
                        if has_media: break
                    if has_media: break
                
                if has_media: continue 
                
                protect_table = False
                if is_ym_mode:
                    try:
                        tbl_text_upper = ""
                        for row in table.rows:
                            for cell in row.cells:
                                tbl_text_upper += cell.text.upper() + " "
                        if ("TRUE" in tbl_text_upper and "FALSE" in tbl_text_upper) or \
                           ("YES" in tbl_text_upper and "NO" in tbl_text_upper) or \
                           ("NOT GIVEN" in tbl_text_upper) or \
                           ("T" in tbl_text_upper and "F" in tbl_text_upper and "QUESTION" in tbl_text_upper):
                            protect_table = True
                    except Exception: pass
                
                if protect_table: continue
                
                is_layout = False
                checked_cells = 0
                protect_table_strong = False 
                
                for row in table.rows:
                    for cell in row.cells:
                        checked_cells += 1
                        for p in cell.paragraphs:
                            txt = p.text.strip().replace('\xa0', ' ').replace('\u200b', '')
                            if txt:
                                if re.match(r'^(thí nghiệm|hiện tượng|chất|thuốc thử|mẫu thử|kết quả|nhiệt độ|dung dịch|tên|công thức)', txt, re.IGNORECASE):
                                    protect_table_strong = True
                                    break
                                if re.match(r'^[\s]*(?:Câu|Question)\s*\d+[\.\:\s]', txt, re.IGNORECASE) or \
                                   re.match(r'^[\s]*#?[A-Da-d]\s*[\.\)]', txt) or \
                                   re.match(r'^[\s\W]*PHẦN\s+([IVX]+)', txt, re.IGNORECASE):
                                    is_layout = True
                                break 
                        if protect_table_strong or is_layout or checked_cells >= 4: break 
                    if protect_table_strong or is_layout or checked_cells >= 4: break
                
                if protect_table_strong: continue 
                
                if is_layout:
                    for row in table.rows:
                        for cell in row.cells:
                            cells_elems = []
                            for child in cell._element: cells_elems.append(child)
                            for child in cells_elems:
                                if child.tag in (qn('w:p'), qn('w:tbl')):
                                    tbl.addprevious(copy.deepcopy(child))
                    if tbl.getparent() is not None:
                        tbl.getparent().remove(tbl)
                    changed = True
                    break

    def convert_floating_to_inline(self, doc_obj):
        body = doc_obj._body._body
        dist_emu = '144000'
        
        anchors = []
        for anchor in body.xpath('.//*[local-name()="anchor"]'): anchors.append(anchor)
        
        for anchor in anchors:
            align = 'left'
            posH_list = anchor.xpath('.//*[local-name()="positionH"]')
            if posH_list:
                align_nodes = posH_list[0].xpath('.//*[local-name()="align"]')
                if align_nodes and align_nodes[0].text in ('left', 'center', 'right'):
                    align = align_nodes[0].text
                    
            inline = OxmlElement('wp:inline')
            inline.set('distT', dist_emu)
            inline.set('distB', dist_emu)
            inline.set('distL', dist_emu)
            inline.set('distR', dist_emu)
            
            anchor_children = []
            for child in anchor: anchor_children.append(child)
            
            for child in anchor_children:
                tag_str = str(child.tag)
                if tag_str.endswith('}extent') or tag_str.endswith('}effectExtent') or tag_str.endswith('}docPr') or tag_str.endswith('}cNvGraphicFramePr') or tag_str.endswith('}graphic'):
                    inline.append(copy.deepcopy(child))
                    
            drawing_node = anchor.getparent()
            if drawing_node is not None and str(drawing_node.tag).endswith('drawing'):
                parent_p = drawing_node.getparent()
                while parent_p is not None and str(parent_p.tag) != qn('w:p'):
                    parent_p = parent_p.getparent()
                    
                if parent_p is not None:
                    new_p = OxmlElement('w:p')
                    pPr = OxmlElement('w:pPr')
                    jc = OxmlElement('w:jc')
                    jc.set(qn('w:val'), align)
                    pPr.append(jc)
                    new_p.append(pPr)
                    
                    new_r = OxmlElement('w:r')
                    new_drawing = OxmlElement('w:drawing')
                    new_drawing.append(inline)
                    new_r.append(new_drawing)
                    new_p.append(new_r)
                    
                    parent_p.addnext(new_p)
                    drawing_node.getparent().remove(drawing_node)
                else:
                    anchor.getparent().replace(anchor, inline)
            else:
                anchor.getparent().replace(anchor, inline)
                
        inlines = []
        for inline_node in body.xpath('.//*[local-name()="inline"]'): inlines.append(inline_node)
        for inline_node in inlines:
            inline_node.set('distT', dist_emu)
            inline_node.set('distB', dist_emu)
            inline_node.set('distL', dist_emu)
            inline_node.set('distR', dist_emu)

        shapes = []
        for shape in body.xpath('.//*[local-name()="shape"]'): shapes.append(shape)
        for shape in shapes:
            style = shape.get('style', '')
            if style:
                style = re.sub(r'position\s*:\s*absolute\s*;?', '', style)
                style = re.sub(r'z-index\s*:\s*[^;]+;?', '', style)
                style = re.sub(r'margin-top\s*:\s*[^;]+;?', '', style)
                style = re.sub(r'margin-left\s*:\s*[^;]+;?', '', style)
                shape.set('style', style)

    def apply_global_font(self, doc):
        try:
            doc.styles['Normal'].font.name = 'Times New Roman'
            doc.styles['Normal'].font.size = Pt(12)
        except: pass

        runs = []
        for r in doc._body._body.xpath('.//w:r'): runs.append(r)
        
        for r in runs:
            rPr = r.find(qn('w:rPr'))
            if rPr is None:
                rPr = OxmlElement('w:rPr')
                r.insert(0, rPr)
                
            rFonts = rPr.find(qn('w:rFonts'))
            if rFonts is None:
                rFonts = OxmlElement('w:rFonts')
                rPr.append(rFonts)
            rFonts.set(qn('w:ascii'), 'Times New Roman')
            rFonts.set(qn('w:hAnsi'), 'Times New Roman')
            rFonts.set(qn('w:eastAsia'), 'Times New Roman')
            rFonts.set(qn('w:cs'), 'Times New Roman')
            
            for attr in ('asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme'):
                if qn(f'w:{attr}') in rFonts.attrib:
                    del rFonts.attrib[qn(f'w:{attr}')]
                    
            rStyle = rPr.find(qn('w:rStyle'))
            if rStyle is not None:
                rPr.remove(rStyle)
            
            for attr in ('themeAscii', 'themeHAnsi', 'themeEastAsia', 'themeCS'):
                if qn(f'w:{attr}') in rFonts.attrib:
                    del rFonts.attrib[qn(f'w:{attr}')]

            color = rPr.find(qn('w:color'))
            if color is None:
                color = OxmlElement('w:color')
                color.set(qn('w:val'), '000000')
                rPr.append(color)
            else:
                if color.get(qn('w:val')) != 'FF0000':
                    color.set(qn('w:val'), '000000')

            for attr in ('themeColor', 'themeTint', 'themeShade'):
                if color is not None and qn(f'w:{attr}') in color.attrib:
                    del color.attrib[qn(f'w:{attr}')]

            t_texts = [t.text for t in r.iter(qn('w:t')) if t.text]
            run_text = "".join(t_texts)
            
            kythi_text = self.val_kythi.get().upper()
            kt_hd_text = kythi_text
            if "KIỂM TRA" in kt_hd_text:
                kt_hd_text = kt_hd_text.replace("KIỂM TRA", "HƯỚNG DẪN CHẤM KIỂM TRA", 1)
            else:
                kt_hd_text = "HƯỚNG DẪN CHẤM " + kt_hd_text

            is_title = False
            if (kythi_text in run_text and kythi_text != "") or (kt_hd_text in run_text and kt_hd_text != ""):
                is_title = True

            sz = rPr.find(qn('w:sz'))
            if sz is None:
                sz = OxmlElement('w:sz')
                rPr.append(sz)
                
            szCs = rPr.find(qn('w:szCs'))
            if szCs is None:
                szCs = OxmlElement('w:szCs')
                rPr.append(szCs)

            if is_title:
                sz.set(qn('w:val'), '26')
                szCs.set(qn('w:val'), '26')
            else:
                sz.set(qn('w:val'), '24')
                szCs.set(qn('w:val'), '24')

            has_obj = False
            for n in r.iter():
                tag = str(n.tag)
                if tag.endswith('}object') or tag.endswith('}drawing') or tag.endswith('}shape') or tag.endswith('}pict') or tag.endswith('}oMath') or tag.endswith('}AlternateContent') or tag.endswith('}OLEObject'):
                    has_obj = True
                    break
            if not has_obj:
                pos = rPr.find(qn('w:position'))
                if pos is not None:
                    rPr.remove(pos)

        ps = []
        for p in doc._body._body.xpath('.//w:p'): ps.append(p)
        for p in ps:
            pPr = p.find(qn('w:pPr'))
            if pPr is not None:
                textAlign = pPr.find(qn('w:textAlignment'))
                if textAlign is not None:
                    pPr.remove(textAlign)

    def force_bold_label(self, p_elm, label_text):
        target_p = p_elm
        if p_elm.tag == qn('w:tbl'):
            for p in p_elm.xpath('.//w:p'):
                if "".join([t.text for t in p.xpath('.//w:t') if t.text]).strip():
                    target_p = p
                    break
            if target_p.tag != qn('w:p'): return p_elm
        elif p_elm.tag != qn('w:p'): 
            return p_elm

        target_p = self.clean_empty_lines(target_p)
            
        t_texts = []
        for t in target_p.iter(qn('w:t')):
            if t.text: t_texts.append(t.text)
        full_text = "".join(t_texts)
        
        pattern = re.compile(r'^\s*(#?(?:Câu|Question)\s*\d+\s*(\([^)]+\))?[\.\:\s]*|#?[A-Da-d]\s*[\.\)]\s*)', re.IGNORECASE)
        match = pattern.match(full_text)
        if match:
            to_del = len(match.group(0))
            for t in target_p.iter(qn('w:t')):
                if not t.text: continue
                if to_del <= 0: break
                l = len(t.text)
                if l <= to_del: 
                    t.text = ""
                    to_del -= l
                else: 
                    t.text = t.text[to_del:].lstrip()
                    to_del = 0
                    
        target_p = self.clean_empty_lines(target_p)
        for t in target_p.iter(qn('w:t')):
            if t.text and t.text.strip(): 
                t.text = t.text.lstrip()
                if t.text.startswith('#'):
                    t.text = t.text[1:].lstrip()
                break
                
        texts = []
        for t in target_p.iter(qn('w:t')):
            if t.text: texts.append(t)
        if texts: texts[-1].text = texts[-1].text.rstrip()
        
        label_text = label_text.replace('#', '')
        
        new_r = OxmlElement('w:r')
        rPr = OxmlElement('w:rPr')
        
        rFonts = OxmlElement('w:rFonts')
        rFonts.set(qn('w:ascii'), 'Times New Roman')
        rFonts.set(qn('w:hAnsi'), 'Times New Roman')
        rFonts.set(qn('w:eastAsia'), 'Times New Roman')
        rFonts.set(qn('w:cs'), 'Times New Roman')
        sz = OxmlElement('w:sz'); sz.set(qn('w:val'), '24')
        szCs = OxmlElement('w:szCs'); szCs.set(qn('w:val'), '24')
        color = OxmlElement('w:color'); color.set(qn('w:val'), '000000')
        b = OxmlElement('w:b'); b.set(qn('w:val'), '1')
        bCs = OxmlElement('w:bCs'); bCs.set(qn('w:val'), '1')
        
        rPr.append(rFonts)
        rPr.append(sz)
        rPr.append(szCs)
        rPr.append(color)
        rPr.append(b)
        rPr.append(bCs)
        new_r.append(rPr)
        
        t = OxmlElement('w:t')
        t.text = label_text.strip() + " "
        t.set(qn('xml:space'), 'preserve')
        new_r.append(t)
        
        pPr = target_p.find(qn('w:pPr'))
        if pPr is None: 
            pPr = OxmlElement('w:pPr')
            target_p.insert(0, pPr)
        pPr.addnext(new_r)
        
        return p_elm

    def get_option_data(self, ans_obj, label):
        base_p = copy.deepcopy(ans_obj["elements"][0])
        extra_elements = []
        for i in range(1, len(ans_obj["elements"])):
            extra_elements.append(ans_obj["elements"][i])
            
        blocks_to_return = [base_p]
        
        for extra_p in extra_elements:
            if extra_p.tag == qn('w:tbl'):
                blocks_to_return.append(copy.deepcopy(extra_p))
            else:
                space_r = OxmlElement('w:r')
                space_t = OxmlElement('w:t')
                space_t.text = " "
                space_t.set(qn('xml:space'), 'preserve')
                space_r.append(space_t)
                base_p.append(space_r)
                for child in extra_p:
                    if str(child.tag) != qn('w:pPr'): base_p.append(copy.deepcopy(child))
        
        base_p = self.unwrap_omathpara(base_p)
        
        brs = []
        for br in base_p.iter(qn('w:br')): brs.append(br)
        for br in brs:
            if br.getparent() is not None: br.getparent().remove(br)
            
        tabs = []
        for tab in base_p.iter(qn('w:tab')): tabs.append(tab)
        for tab in tabs:
            if tab.getparent() is not None: tab.getparent().remove(tab)
                    
        p_ans = self.force_bold_label(base_p, label)
        if p_ans.tag == qn('w:p'):
            p_ans = self.clean_empty_lines(p_ans)
            
        txt = ""
        for node in p_ans.xpath('.//w:t | .//m:t'):
            if node.text: txt += node.text
            
        txt_calc = re.sub(r'\s+', ' ', txt).strip()
        
        math_complex = 0
        for n in p_ans.xpath('.//*[local-name()="f"] | .//*[local-name()="rad"] | .//*[local-name()="nary"] | .//*[local-name()="sSup"] | .//*[local-name()="sSub"] | .//*[local-name()="sSubSup"]'):
            math_complex += 0.2

        math_small = 0
        for n in p_ans.xpath('.//*[local-name()="oMath"]'):
            math_small += 0.5
        
        total_width_cm = len(txt_calc) * 0.22 + math_small + math_complex

        has_ole = False
        has_image = False
        
        for container in p_ans.xpath('.//*[local-name()="object"] | .//*[local-name()="drawing"] | .//*[local-name()="pict"]'):
            has_image = True
            if container.tag.endswith('}object'):
                has_ole = True
                
            container_w_cm = 0.5
            extents = container.xpath('.//*[local-name()="extent"] | .//*[local-name()="ext"]')
            if extents:
                cx = int(extents[0].get('cx', '0'))
                if cx > 0:
                    container_w_cm += cx / 360000
            else:
                shapes = container.xpath('.//*[local-name()="shape"]')
                if shapes:
                    style = shapes[0].get('style', '')
                    w_match = re.search(r'width:\s*([0-9\.]+)\s*(pt|cm|in)', style, re.IGNORECASE)
                    if w_match:
                        w = float(w_match.group(1))
                        unit = w_match.group(2).lower()
                        if unit == 'cm': container_w_cm += w
                        elif unit == 'in': container_w_cm += w * 2.54
                        else: container_w_cm += w / 28.3465
                    else:
                        container_w_cm += 1.0
            total_width_cm += container_w_cm

        if has_image and not has_ole and len(txt_calc) <= 10:
            total_width_cm = min(total_width_cm, 3.8)

        return blocks_to_return, total_width_cm

    def scale_images_in_paragraph(self, p_elm, max_width_cm):
        max_width_emu = int(max_width_cm * 360000)
        for node in p_elm.iter():
            tag = str(node.tag)
            if tag.endswith('}extent') or tag.endswith('}ext'):
                cx = int(node.get('cx', '0'))
                cy = int(node.get('cy', '0'))
                if cx > max_width_emu and cx > 0:
                    ratio = max_width_emu / cx
                    node.set('cx', str(max_width_emu))
                    node.set('cy', str(int(cy * ratio)))
                    
            elif tag.endswith('}shape'):
                style = node.get('style', '')
                if style:
                    w_match = re.search(r'width:([0-9\.]+)pt', style)
                    h_match = re.search(r'height:([0-9\.]+)pt', style)
                    if w_match and h_match:
                        w = float(w_match.group(1))
                        h = float(h_match.group(1))
                        max_width_pt = max_width_cm * 28.3465
                        if w > max_width_pt:
                            ratio = max_width_pt / w
                            new_w = w * ratio
                            new_h = h * ratio
                            new_style = re.sub(r'width:[0-9\.]+pt', f'width:{new_w:.2f}pt', style)
                            new_style = re.sub(r'height:[0-9\.]+pt', f'height:{new_h:.2f}pt', new_style)
                            node.set('style', new_style)

    def append_options_as_text(self, doc, body, opts_ready, cols, indent_cm=0.5):
        if cols == 1:
            for opt_tuple in opts_ready:
                blocks_list = opt_tuple[0]
                for b_idx, blk in enumerate(blocks_list):
                    if blk.tag == qn('w:p'):
                        if b_idx == 0:
                            self.reset_paragraph_indent(blk, left_cm=indent_cm, first_line_cm=0)
                        self.force_zero_spacing(blk, justify=False, preserve_align=False)
                    safe_append(body, blk)
            return
            
        opts_list = []
        for i in range(len(opts_ready)): opts_list.append(opts_ready[i])
        rows = (len(opts_list) + cols - 1) // cols
        
        page_width_cm = 18.0 - indent_cm
        col_width_cm = page_width_cm / cols
        
        for r in range(rows):
            new_p = OxmlElement('w:p')
            pPr = OxmlElement('w:pPr')
            
            tabs = OxmlElement('w:tabs')
            for c in range(1, cols):
                tab = OxmlElement('w:tab')
                tab.set(qn('w:val'), 'left')
                tab.set(qn('w:pos'), str(int((indent_cm + c * col_width_cm) * 567)))
                tabs.append(tab)
            pPr.append(tabs)
            new_p.append(pPr)
            
            self.reset_paragraph_indent(new_p, left_cm=indent_cm, first_line_cm=0)
            self.force_zero_spacing(new_p, justify=False, preserve_align=False)
            
            for c in range(cols):
                idx = r * cols + c
                if idx < len(opts_list):
                    blocks_list = opts_list[idx][0]
                    first_p = blocks_list[0] 
                    
                    if c > 0:
                        r_tab = OxmlElement('w:r')
                        r_tab.append(OxmlElement('w:tab'))
                        new_p.append(r_tab)
                        
                    self.scale_images_in_paragraph(first_p, col_width_cm - 0.2)
                    
                    children_p_opt = []
                    for child in first_p: children_p_opt.append(child)
                    for child in children_p_opt:
                        if str(child.tag) != qn('w:pPr'):
                            new_p.append(copy.deepcopy(child))
            
            safe_append(body, new_p)
            
            for c in range(cols):
                idx = r * cols + c
                if idx < len(opts_list):
                    blocks_list = opts_list[idx][0]
                    if len(blocks_list) > 1:
                        for extra_blk in blocks_list[1:]:
                            safe_append(body, copy.deepcopy(extra_blk))

    def remove_youngmix_tags(self, doc_obj):
        for p in doc_obj.paragraphs:
            text_lower = p.text.lower()
            if '<g' in text_lower or '<#g' in text_lower or '＜g' in text_lower or '＜#g' in text_lower:
                for run in p.runs:
                    if run.text:
                        run.text = re.sub(r'[<＜]#?g\d*[>＞]', '', run.text, flags=re.IGNORECASE)
                if not p.text.strip() and not any(str(n.tag).endswith(('drawing','pict','object','oMath')) for n in p._element.iter()):
                    if p._element.getparent() is not None:
                        p._element.getparent().remove(p._element)
                    
        for table in doc_obj.tables:
            for row in table.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        text_lower = p.text.lower()
                        if '<g' in text_lower or '<#g' in text_lower:
                            for run in p.runs:
                                if run.text:
                                    run.text = re.sub(r'<#?g\d*>', '', run.text, flags=re.IGNORECASE)
                            if not p.text.strip() and not any(str(n.tag).endswith(('drawing','pict','object','oMath')) for n in p._element.iter()):
                                if p._element.getparent() is not None:
                                    p._element.getparent().remove(p._element)

    def clear_answers_from_exam(self, doc_obj):
        for rPr in doc_obj._body._body.xpath('.//w:rPr'):
            for tag in ('w:color', 'w:shd', 'w:highlight'):
                elem = rPr.find(qn(tag))
                if elem is not None:
                    rPr.remove(elem)

    def _setup_footer(self, ftr, ma, p_type="NUMPAGES", is_english_exam=False):
        ftr.is_linked_to_previous = False
        for child in list(ftr._element):
            ftr._element.remove(child)
            
        p0 = ftr.add_paragraph()
        p0.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p0.paragraph_format.tab_stops.add_tab_stop(Cm(18.0), WD_TAB_ALIGNMENT.RIGHT)
        
        show_ma = ma
        if ma == "GOC": show_ma = "GỐC"
        run_l = p0.add_run(f"Mã đề: {show_ma}")
        run_l.font.name = "Times New Roman"
        run_l.font.size = Pt(12)

        run_t = p0.add_run("\tTrang ")
        run_t.font.name = "Times New Roman"
        run_t.font.size = Pt(12)
        
        run_pg = p0.add_run()
        fld1 = OxmlElement('w:fldChar'); fld1.set(qn('w:fldCharType'), 'begin')
        instr1 = OxmlElement('w:instrText'); instr1.set(qn('xml:space'), 'preserve'); instr1.text = " PAGE "
        fld2 = OxmlElement('w:fldChar'); fld2.set(qn('w:fldCharType'), 'separate')
        t1 = OxmlElement('w:t'); t1.text = "1"
        fld3 = OxmlElement('w:fldChar'); fld3.set(qn('w:fldCharType'), 'end')
        
        run_pg._r.append(fld1)
        run_pg._r.append(instr1)
        run_pg._r.append(fld2)
        run_pg._r.append(t1)
        run_pg._r.append(fld3)

        run_slash = p0.add_run("/")
        run_slash.font.name = "Times New Roman"
        run_slash.font.size = Pt(12)

        run_num = p0.add_run()
        fld4 = OxmlElement('w:fldChar'); fld4.set(qn('w:fldCharType'), 'begin')
        instr2 = OxmlElement('w:instrText'); instr2.set(qn('xml:space'), 'preserve'); instr2.text = f" {p_type} "
        fld5 = OxmlElement('w:fldChar'); fld5.set(qn('w:fldCharType'), 'separate')
        t2 = OxmlElement('w:t'); t2.text = "..."
        fld6 = OxmlElement('w:fldChar'); fld6.set(qn('w:fldCharType'), 'end')
        
        run_num._r.append(fld4)
        run_num._r.append(instr2)
        run_num._r.append(fld5)
        run_num._r.append(t2)
        run_num._r.append(fld6)

    def generate_header_elements(self, ma_de, is_english_exam=False):
        td = docx.Document()
        tb = td.add_table(rows=1, cols=2)
        tb.autofit = False
        tb.alignment = WD_TABLE_ALIGNMENT.CENTER
        
        col_idx = 0
        for col in tb.columns:
            if col_idx == 0: col.width = Cm(13.5)
            elif col_idx == 1: col.width = Cm(4.5)
            col_idx += 1
        
        tp = tb._element.find(qn('w:tblPr'))
        if tp is not None:
            bds = OxmlElement('w:tblBorders')
            bd_top = OxmlElement('w:top'); bd_top.set(qn('w:val'), 'none'); bds.append(bd_top)
            bd_left = OxmlElement('w:left'); bd_left.set(qn('w:val'), 'none'); bds.append(bd_left)
            bd_bottom = OxmlElement('w:bottom'); bd_bottom.set(qn('w:val'), 'none'); bds.append(bd_bottom)
            bd_right = OxmlElement('w:right'); bd_right.set(qn('w:val'), 'none'); bds.append(bd_right)
            bd_inH = OxmlElement('w:insideH'); bd_inH.set(qn('w:val'), 'none'); bds.append(bd_inH)
            bd_inV = OxmlElement('w:insideV'); bd_inV.set(qn('w:val'), 'none'); bds.append(bd_inV)
            tp.append(bds)

        c0 = tb.cell(0, 0)
        p0 = c0.paragraphs[0]
        self.force_zero_spacing(p0._element, False, False)
        p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self.reset_paragraph_indent(p0._element, 0, 0)
        
        sotx = self.val_so.get().upper().replace("SỞ GIÁO DỤC VÀ ĐÀO TẠO ", "SỞ GIÁO DỤC VÀ ĐÀO TẠO\n")
        r01 = p0.add_run(sotx + "\n")
        r01.font.name = "Times New Roman"
        r01.font.size = Pt(12)
        
        r02 = p0.add_run(self.val_truong.get().upper() + "\n")
        r02.font.name = "Times New Roman"
        r02.font.size = Pt(12)
        r02.font.bold = True
        
        r03 = p0.add_run("-------***-------")
        r03.font.name = "Times New Roman"
        r03.font.size = Pt(12)
        
        ppg = c0.add_paragraph()
        self.force_zero_spacing(ppg._element, False, False)
        ppg.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self.reset_paragraph_indent(ppg._element, 0, 0)
        
        r_pg1 = ppg.add_run("(Đề có . trang)")
        r_pg1.font.name = "Times New Roman"
        r_pg1.font.size = Pt(12)
        r_pg1.font.italic = True

        c1 = tb.cell(0, 1)
        p1 = c1.paragraphs[0]
        self.force_zero_spacing(p1._element, False, False)
        p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self.reset_paragraph_indent(p1._element, 0, 0)
        
        r11 = p1.add_run(self.val_kythi.get().upper() + "\n")
        r11.font.name = "Times New Roman"
        r11.font.size = Pt(13) 
        r11.font.bold = True
        
        r12 = p1.add_run(self.val_namhoc.get().upper() + "\n")
        r12.font.name = "Times New Roman"
        r12.font.size = Pt(12)
        r12.font.bold = True
        
        r13 = p1.add_run(self.val_mon.get() + "\n")
        r13.font.name = "Times New Roman"
        r13.font.size = Pt(12)
        r13.font.bold = True
        
        r14 = p1.add_run(self.val_thoigian.get())
        r14.font.name = "Times New Roman"
        r14.font.size = Pt(12)
        r14.font.italic = True
        
        pe = td.add_paragraph()
        self.force_zero_spacing(pe._element, False, False)
        r_emp = pe.add_run(" ")
        r_emp.font.name = "Times New Roman"
        r_emp.font.size = Pt(12)
        pe._element.set(qn('w:rsidR'), 'KEEP_ME_EMPTY_LINE') 

        pi = td.add_paragraph()
        pPr = pi._element.get_or_add_pPr()
        pStyle = pPr.find(qn('w:pStyle'))
        if pStyle is None:
            pStyle = OxmlElement('w:pStyle')
            pPr.insert(0, pStyle)
        pStyle.set(qn('w:val'), 'NoSpacing')

        spacing = pPr.find(qn('w:spacing'))
        if spacing is not None:
            pPr.remove(spacing)

        pi.alignment = WD_ALIGN_PARAGRAPH.LEFT
        self.reset_paragraph_indent(pi._element, 0, 0)
        pi.paragraph_format.tab_stops.clear_all()
        pi.paragraph_format.tab_stops.add_tab_stop(Cm(17.0), WD_TAB_ALIGNMENT.RIGHT)
        
        ri = pi.add_run("Họ và tên...................................................Lớp..............; Số báo danh....................\t")
        ri.font.name = "Times New Roman"
        ri.font.size = Pt(12)
        
        display_ma = "GỐC"
        if ma_de != "GOC": display_ma = ma_de
        
        rma = pi.add_run(f"Mã đề: {display_ma}")
        rma.font.name = "Times New Roman"
        rma.font.size = Pt(12)
        rma.font.bold = True
        
        rprma = rma._r.get_or_add_rPr()
        bdr = OxmlElement('w:bdr')
        bdr.set(qn('w:val'), 'single')
        bdr.set(qn('w:sz'), '8')
        bdr.set(qn('w:space'), '4') 
        bdr.set(qn('w:color'), '000000')
        rprma.append(bdr)
        
        result_elements = []
        for elm in td._body._body:
            if elm.tag != qn('w:sectPr'):
                result_elements.append(copy.deepcopy(elm))
        return result_elements

    def add_hd_cham_header(self, doc_obj):
        tb_hd = doc_obj.add_table(rows=1, cols=2)
        tb_hd.autofit = False
        tb_hd.alignment = WD_TABLE_ALIGNMENT.CENTER
        
        col_idx = 0
        for col in tb_hd.columns:
            if col_idx == 0: col.width = Cm(8.5)
            elif col_idx == 1: col.width = Cm(9.5)
            col_idx += 1
        
        tp_hd = tb_hd._element.find(qn('w:tblPr'))
        if tp_hd is not None:
            bds_hd = OxmlElement('w:tblBorders')
            bd_top = OxmlElement('w:top'); bd_top.set(qn('w:val'), 'none'); bds_hd.append(bd_top)
            bd_left = OxmlElement('w:left'); bd_left.set(qn('w:val'), 'none'); bds_hd.append(bd_left)
            bd_bottom = OxmlElement('w:bottom'); bd_bottom.set(qn('w:val'), 'none'); bds_hd.append(bd_bottom)
            bd_right = OxmlElement('w:right'); bd_right.set(qn('w:val'), 'none'); bds_hd.append(bd_right)
            bd_inH = OxmlElement('w:insideH'); bd_inH.set(qn('w:val'), 'none'); bds_hd.append(bd_inH)
            bd_inV = OxmlElement('w:insideV'); bd_inV.set(qn('w:val'), 'none'); bds_hd.append(bd_inV)
            tp_hd.append(bds_hd)

        c0_hd = tb_hd.cell(0, 0)
        c0_hd.text = "" 
        
        p0_hd = None
        if c0_hd.paragraphs: p0_hd = c0_hd.paragraphs[0]
        else: p0_hd = c0_hd.add_paragraph()
        
        self.force_zero_spacing(p0_hd._element, False, False)
        p0_hd.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self.reset_paragraph_indent(p0_hd._element, 0, 0)
        
        sotx_hd = self.val_so.get().upper().replace("SỞ GIÁO DỤC VÀ ĐÀO TẠO ", "SỞ GIÁO DỤC VÀ ĐÀO TẠO\n")
        r01_hd = p0_hd.add_run(sotx_hd + "\n")
        r01_hd.font.name = "Times New Roman"
        r01_hd.font.size = Pt(12)
        r02_hd = p0_hd.add_run(self.val_truong.get().upper() + "\n")
        r02_hd.font.name = "Times New Roman"
        r02_hd.font.size = Pt(12)
        r02_hd.font.bold = True
        r03_hd = p0_hd.add_run("-------***-------")
        r03_hd.font.name = "Times New Roman"
        r03_hd.font.size = Pt(12)

        c1_hd = tb_hd.cell(0, 1)
        c1_hd.text = "" 
        
        p1_hd = None
        if c1_hd.paragraphs: p1_hd = c1_hd.paragraphs[0]
        else: p1_hd = c1_hd.add_paragraph()
        
        self.force_zero_spacing(p1_hd._element, False, False)
        
        p1_hd.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self.reset_paragraph_indent(p1_hd._element, 0, 0)
        
        kt_tx_hd = self.val_kythi.get().upper()
        if "KIỂM TRA" in kt_tx_hd:
            kt_tx_hd = kt_tx_hd.replace("KIỂM TRA", "HƯỚNG DẪN CHẤM KIỂM TRA", 1)
        else:
            kt_tx_hd = "HƯỚNG DẪN CHẤM " + kt_tx_hd
            
        r11_hd = p1_hd.add_run(kt_tx_hd + "\n")
        
        r11_hd.font.name = "Times New Roman"
        r11_hd.font.size = Pt(13)
        r11_hd.font.bold = True
        
        r12_hd = p1_hd.add_run(self.val_namhoc.get().upper() + "\n")
        r12_hd.font.name = "Times New Roman"
        r12_hd.font.size = Pt(12)
        r12_hd.font.bold = True
        r13_hd = p1_hd.add_run(self.val_mon.get())
        r13_hd.font.name = "Times New Roman"
        r13_hd.font.size = Pt(12)
        r13_hd.font.bold = True
        
        pe = doc_obj.add_paragraph()
        self.force_zero_spacing(pe._element, False, False)

    def update_table_questions(self, element, current_q_idx):
        is_ym = self.youngmix_var.get()
        if not is_ym:
            return element, current_q_idx
            
        if element.tag == qn('w:tbl'):
            tmp_doc = docx.Document()
            tmp_doc._body._body.append(element)
            tmp_tbl = tmp_doc.tables[0]
            q_re = re.compile(r'^\s*(Câu|Question)\s*\d+', re.IGNORECASE)
            for row in tmp_tbl.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        m_q = q_re.match(p.text.strip().replace('\xa0', ' ').replace('\u200b', ''))
                        if m_q:
                            q_word = m_q.group(1).capitalize()
                            self.force_bold_label(p._element, f"{q_word} {current_q_idx}. ")
                            current_q_idx += 1
            return tmp_tbl._element, current_q_idx
        return element, current_q_idx

    def process_web(self):
        try:
            with open(self.filepath, "rb") as f:
                self.file_bytes_in_ram = f.read()

            self.global_doc = docx.Document(io.BytesIO(self.file_bytes_in_ram))
            self.convert_auto_numbering(self.global_doc)
            self.convert_floating_to_inline(self.global_doc)
            self.split_soft_returns(self.global_doc)
            self.preprocess_textboxes(self.global_doc)
            self.preprocess_layout_tables(self.global_doc)
            
            is_ym = self.youngmix_var.get()
            if is_ym:
                self.parsed_data, self.errors = self.parse_document_structure(self.global_doc)
                self.parse_youngmix_structure(self.global_doc)
            else:
                self.parsed_data, self.errors = self.parse_document_structure(self.global_doc)

            raw_ma_text = self.ma_de_str
            
            raw_ma = raw_ma_text.split(",")
            ma_des = []
            for m in raw_ma:
                if m.strip(): ma_des.append(m.strip())
                
            if not ma_des:
                raise Exception("Vui lòng nhập ít nhất 1 mã đề để trộn!")
                
            base_dir = self.output_dir if self.output_dir else os.path.dirname(self.filepath)
            
            mon_name_raw = self.val_mon.get()
            mon_name = mon_name_raw.replace("Môn:", "").replace("Môn", "").strip()
            safe_mon_name = re.sub(r'[\\/*?:"<>|]', "", mon_name)
            if not safe_mon_name: safe_mon_name = "De_Tron"
            out_dir = os.path.join(base_dir, f"{safe_mon_name}_{random.randint(100,999)}")
            os.makedirs(out_dir, exist_ok=True)
            
            ma_des_full = ["GOC"]
            for ma in ma_des: ma_des_full.append(ma)
            
            answers_data = {}
            for ma in ma_des_full:
                answers_data[ma] = {'P1': [], 'P2': [], 'P3': []}
                
            prev_exam_orders = {}
            prev_selected_qs = {} 

            doc_all = docx.Document(io.BytesIO(self.file_bytes_in_ram))
            body_all = doc_all._body._body
            
            children_body_all = []
            for child in body_all: children_body_all.append(child)
            for child in children_body_all:
                if child.tag != qn('w:sectPr'):
                    body_all.remove(child)
                    
            is_first_for_all = True
            
            is_english_exam = False
            for p_chk in self.parsed_data.get("parts", []):
                for m_chk in p_chk.get("mucs", []):
                    for q_chk in m_chk.get("questions", []):
                        if "question" in q_chk.get('raw_text', '').lower():
                            is_english_exam = True
                            break
                    if is_english_exam: break
                if is_english_exam: break
            
            total_tasks = len(ma_des_full) + 1 
            current_task = 0

            for ma in ma_des_full:
                current_task += 1
                doc_de = docx.Document(io.BytesIO(self.file_bytes_in_ram))
                self.convert_auto_numbering(doc_de)
                
                for section in doc_de.sections:
                    section.page_width = Cm(21.0)
                    section.page_height = Cm(29.7)
                    section.top_margin = Cm(1.5)
                    section.bottom_margin = Cm(1.5)
                    section.left_margin = Cm(1.5)
                    section.right_margin = Cm(1.5)
                    section.header_distance = Cm(0.8) 
                    section.footer_distance = Cm(0.6) 
                    
                    self._setup_footer(section.footer, ma, "NUMPAGES", is_english_exam)
                    self._setup_footer(section.first_page_footer, ma, "NUMPAGES", is_english_exam)
                
                body = doc_de._body._body
                children_body = []
                for child in body: children_body.append(child)
                for child in children_body:
                    if child.tag != qn('w:sectPr'):
                        body.remove(child)
                    
                p1_ans = []
                p2_ans = []
                p3_ans = [] 
                
                header_elements = self.generate_header_elements(ma, is_english_exam)
                for el in header_elements: safe_append(body, el)
                
                if len(header_elements) > 0:
                    p_info_last = header_elements[-1]
                    if p_info_last.tag == qn('w:p'):
                        p_last = Paragraph(p_info_last, doc_de)
                        try:
                            p_last.style = doc_de.styles['No Spacing']
                        except: pass
                        pPr = p_info_last.get_or_add_pPr()
                        spc = pPr.find(qn('w:spacing'))
                        if spc is not None: pPr.remove(spc)
                
                q_idx = 1
                
                for p_idx, part in enumerate(self.parsed_data.get("parts", [])):
                    is_continuous = self.continuous_numbering_var.get()
                    
                    if not is_ym:
                        q_idx = 1 
                    elif is_ym and not is_continuous:
                        q_idx = 1 
                    
                    if part.get("is_explicit", True) and part.get("title_element") is not None:
                        part_title_copy = copy.deepcopy(part["title_element"])
                        if part_title_copy.tag == qn('w:p'):
                            part_title_copy = self.clean_empty_lines(part_title_copy)
                            self.force_zero_spacing(part_title_copy, justify=False, preserve_align=True)
                            self.reset_paragraph_indent(part_title_copy, left_cm=0, first_line_cm=0)
                        pr_t = Paragraph(part_title_copy, doc_de)
                        pr_t.paragraph_format.space_before = Cm(0.3)
                        pr_t.paragraph_format.space_after = Cm(0.1)
                        safe_append(body, part_title_copy)
                    
                    muc_indices = []
                    for i in range(len(part.get("mucs", []))): muc_indices.append(i)
                    
                    fixed_muc_indices = set()
                    for m_idx in muc_indices:
                        ym_ctrl = self.get_ym_ctrl(p_idx, m_idx) if is_ym else None
                        shuf_var_val = True
                        if is_ym and ym_ctrl:
                            is_fixed = ym_ctrl['fix'].get() 
                            shuf_var_val = not is_fixed
                        elif (p_idx, m_idx) in self.muc_order_shuffle_vars:
                            shuf_var_val = self.muc_order_shuffle_vars[(p_idx, m_idx)].get()
                            
                        if not shuf_var_val:
                            fixed_muc_indices.add(m_idx)
                            
                    if ma != "GOC":
                        prev_m_indices = None
                        if f"part_{p_idx}_mucs" in prev_exam_orders:
                            prev_m_indices = prev_exam_orders[f"part_{p_idx}_mucs"]
                        
                        shuffleable_positions = []
                        for i in muc_indices:
                            if i not in fixed_muc_indices: shuffleable_positions.append(i)
                            
                        shuffleable_values = []
                        for i in shuffleable_positions: shuffleable_values.append(muc_indices[i])
                        
                        if prev_m_indices is not None and len(shuffleable_values) > 1:
                            best_shuffle = []
                            for v in shuffleable_values: best_shuffle.append(v)
                            for _ in range(100):
                                random.shuffle(best_shuffle)
                                overlap = False
                                for idx, pos in enumerate(shuffleable_positions):
                                    if best_shuffle[idx] == prev_m_indices[pos]:
                                        overlap = True
                                        break
                                if not overlap:
                                    break
                            shuffleable_values = best_shuffle
                        elif len(shuffleable_values) > 1:
                            random.shuffle(shuffleable_values)
                            
                        final_m_indices = []
                        for v in muc_indices: final_m_indices.append(v)
                        for idx, pos in enumerate(shuffleable_positions):
                            final_m_indices[pos] = shuffleable_values[idx]
                            
                        muc_indices = final_m_indices
                        prev_exam_orders[f"part_{p_idx}_mucs"] = []
                        for v in muc_indices: prev_exam_orders[f"part_{p_idx}_mucs"].append(v)
                    else:
                        prev_exam_orders[f"part_{p_idx}_mucs"] = []
                        for v in muc_indices: prev_exam_orders[f"part_{p_idx}_mucs"].append(v)
                        
                    for output_m_idx, m_idx in enumerate(muc_indices):
                        muc = part["mucs"][m_idx]
                        
                        is_ym_tag_title = False
                        if is_ym:
                            title_text_raw = muc.get("title", "")
                            if title_text_raw and re.match(r'^\s*<#?g\d*>\s*$', title_text_raw.strip(), re.IGNORECASE):
                                is_ym_tag_title = True

                        is_first_in_group = False
                        if is_ym and is_english_exam:
                            is_first_in_group = True

                        if self.keep_muc_var.get() and muc.get("title_element") is not None and not is_ym_tag_title: 
                            muc_title_copy = copy.deepcopy(muc["title_element"])
                            t_texts = []
                            for t in muc_title_copy.iter(qn('w:t')):
                                if t.text: t_texts.append(t.text)
                            full_text = "".join(t_texts)
                            
                            m_muc_prefix = re.match(r'^\s*(?:[IVXLCDM]+|\d+)\s*\.\s*', full_text, re.IGNORECASE)
                            if m_muc_prefix:
                                to_del = len(m_muc_prefix.group(0))
                                for t in muc_title_copy.iter(qn('w:t')):
                                    if not t.text: continue
                                    if to_del <= 0: break
                                    l = len(t.text)
                                    if l <= to_del:
                                        t.text = ""
                                        to_del -= l
                                    else:
                                        t.text = t.text.lstrip()
                                        to_del = 0
                                        
                            if muc_title_copy.tag == qn('w:p'):
                                muc_title_copy = self.clean_empty_lines(muc_title_copy)
                                self.reset_paragraph_indent(muc_title_copy, left_cm=0, first_line_cm=0)
                                self.force_zero_spacing(muc_title_copy, justify=False, preserve_align=True)
                                
                                if is_first_in_group:
                                    Paragraph(muc_title_copy, doc_de).paragraph_format.space_before = Cm(0.3)
                                    is_first_in_group = False
                                    
                            safe_append(body, muc_title_copy)
                        
                        if "description_elements" in muc:
                            for desc_el in muc["description_elements"]:
                                desc_copy = copy.deepcopy(desc_el)
                                if desc_copy.tag == qn('w:p'):
                                    desc_copy = self.clean_empty_lines(desc_copy)
                                    self.reset_paragraph_indent(desc_copy, left_cm=0, first_line_cm=0)
                                    self.force_zero_spacing(desc_copy, justify=False, preserve_align=True)
                                    
                                    if is_first_in_group:
                                        Paragraph(desc_copy, doc_de).paragraph_format.space_before = Cm(0.3)
                                        is_first_in_group = False
                                        
                                elif desc_copy.tag == qn('w:tbl'):
                                    desc_copy, q_idx = self.update_table_questions(desc_copy, q_idx)
                                    is_first_in_group = False 
                                safe_append(body, desc_copy)
                        
                        ym_ctrl = self.get_ym_ctrl(p_idx, m_idx) if is_ym else None
                        
                        if is_ym and ym_ctrl:
                            if ym_ctrl['start_q1'].get():
                                q_idx = 1

                        do_shuffle_questions = True
                        do_shuffle_options = True
                        num_q_user = len(muc.get('questions', []))
                        fix_str = ""

                        if is_ym and ym_ctrl:
                            mix_val = ym_ctrl['mix'].get()
                            do_shuffle_questions = mix_val in ("Chỉ câu hỏi", "Câu hỏi & đáp án")
                            do_shuffle_options = mix_val in ("Chỉ đáp án", "Câu hỏi & đáp án")
                            q_type_val = ym_ctrl['type'].get()
                            if q_type_val == "Tự luận":
                                do_shuffle_options = False
                        
                        total_q_muc = len(muc.get('questions', []))
                        num_q_user = min(num_q_user, total_q_muc)

                        fixed_indices = set()
                        if fix_str.strip():
                            for x in fix_str.split(','):
                                try:
                                    val = int(x.strip()) - 1
                                    if 0 <= val < num_q_user: fixed_indices.add(val)
                                except ValueError: pass
                                
                        for i_q, q_dict in enumerate(muc.get('questions', [])):
                            if q_dict.get('is_fixed'): fixed_indices.add(i_q)

                        valid_fixed_positions = set()
                        for i in fixed_indices:
                            if i < num_q_user: valid_fixed_positions.add(i)

                        if ma == "GOC":
                            final_indices = [i for i in range(num_q_user)]
                            prev_exam_orders[(p_idx, m_idx)] = [i for i in final_indices]
                            prev_selected_qs[(p_idx, m_idx)] = set(final_indices)
                        else:
                            if not do_shuffle_questions:
                                final_indices = [i for i in range(num_q_user)]
                                prev_exam_orders[(p_idx, m_idx)] = [i for i in final_indices]
                                prev_selected_qs[(p_idx, m_idx)] = set(final_indices)
                            else:
                                prev_indices = prev_exam_orders.get((p_idx, m_idx), [])
                                prev_set = prev_selected_qs.get((p_idx, m_idx), set())
                                forced_qs = set(valid_fixed_positions)
                                pool = [i for i in range(total_q_muc) if i not in forced_qs]
                                R = num_q_user - len(forced_qs)
                                fresh_pool = [i for i in pool if i not in prev_set]
                                used_pool = [i for i in pool if i in prev_set]
                                random.shuffle(fresh_pool)
                                random.shuffle(used_pool)
                                picked_qs = fresh_pool + used_pool
                                final_picked_qs = picked_qs[:R]
                                shuffleable_positions = [i for i in range(num_q_user) if i not in valid_fixed_positions]
                                best_shuffle = [x for x in final_picked_qs]
                                if len(prev_indices) > 0 and len(shuffleable_positions) > 1:
                                    for _ in range(100):
                                        random.shuffle(best_shuffle)
                                        overlap = False
                                        for idx, pos in enumerate(shuffleable_positions):
                                            if pos < len(prev_indices) and best_shuffle[idx] == prev_indices[pos]:
                                                overlap = True; break
                                        if not overlap: break
                                final_indices = [0]*num_q_user
                                for p in valid_fixed_positions: final_indices[p] = p
                                for idx, pos in enumerate(shuffleable_positions):
                                    if idx < len(best_shuffle): final_indices[pos] = best_shuffle[idx]
                                prev_exam_orders[(p_idx, m_idx)] = [x for x in final_indices]
                                prev_selected_qs[(p_idx, m_idx)] = set(final_indices)

                        qs = [copy.deepcopy(muc["questions"][i]) for i in final_indices]
                        
                        for q in qs:
                            if q.get('is_virtual'):
                                continue
                            
                            q_type = q.get('ym_type', part["type"]) if is_ym else part["type"]
                            
                            p_q_raw = copy.deepcopy(q['q_elements'][0])
                            p_q_raw = self.unwrap_omathpara(p_q_raw)
                            
                            extra_text = ""
                            m_extract = re.search(r'(?:Câu|Question)\s*\d+\s*(\([^)]+\))?[\.\:\s]', q['raw_text'], re.IGNORECASE)
                            if m_extract and m_extract.group(1):
                                extra_text = m_extract.group(1)
                                    
                            q_prefix = "Câu"
                            if "question" in q['raw_text'].lower():
                                q_prefix = "Question"
                                is_english_exam = True 
                            
                            p_q = self.force_bold_label(p_q_raw, f"{q_prefix} {q_idx}{extra_text}. ")
                            
                            if p_q.tag == qn('w:p'):
                                p_q = self.clean_empty_lines(p_q)
                                self.reset_paragraph_indent(p_q, left_cm=0, first_line_cm=0)
                                self.force_zero_spacing(p_q, justify=False, preserve_align=False)
                                
                                if is_first_in_group:
                                    Paragraph(p_q, doc_de).paragraph_format.space_before = Cm(0.3)
                                    is_first_in_group = False
                            
                            if q_type in (3, 4):
                                short_answer = ""
                                if q_type == 3:
                                    ans = self.extract_and_clean_dapan(p_q)
                                    if ans: short_answer = ans
                                
                                t_texts = []
                                for t in p_q.iter(qn('w:t')):
                                    if t.text: t_texts.append(t.text)
                                txt_pq = "".join(t_texts).strip()
                                
                                has_obj_pq = False
                                for n in p_q.iter():
                                    if str(n.tag).endswith('imagedata') or str(n.tag).endswith('pic') or str(n.tag).endswith('object') or str(n.tag).endswith('drawing') or str(n.tag).endswith('oMath') or str(n.tag).endswith('AlternateContent'):
                                        has_obj_pq = True
                                        break
                                        
                                if txt_pq or has_obj_pq: safe_append(body, p_q)
                                
                                q_elements_skip_first = []
                                for i in range(1, len(q['q_elements'])):
                                    q_elements_skip_first.append(q['q_elements'][i])
                                    
                                for el in q_elements_skip_first: 
                                    ep_q_copy = copy.deepcopy(el)
                                    if ep_q_copy.tag == qn('w:p'):
                                        ep_q_copy = self.unwrap_omathpara(ep_q_copy)
                                        ep_q_copy = self.clean_empty_lines(ep_q_copy)
                                        self.reset_paragraph_indent(ep_q_copy, left_cm=0, first_line_cm=0)
                                        self.force_zero_spacing(ep_q_copy, justify=False, preserve_align=True)
                                    elif ep_q_copy.tag == qn('w:tbl'):
                                        ep_q_copy, q_idx = self.update_table_questions(ep_q_copy, q_idx)
                                        
                                    t_texts_ep = []
                                    for t in ep_q_copy.iter(qn('w:t')):
                                        if t.text: t_texts_ep.append(t.text)
                                    txt_content = "".join(t_texts_ep).strip()
                                    
                                    has_img = False
                                    for n in ep_q_copy.iter():
                                        if str(n.tag).endswith('imagedata') or str(n.tag).endswith('pic') or str(n.tag).endswith('object') or str(n.tag).endswith('drawing') or str(n.tag).endswith('AlternateContent'):
                                            has_img = True
                                            break
                                            
                                    if has_img and not txt_content and ep_q_copy.tag == qn('w:p'):
                                        p_obj = Paragraph(ep_q_copy, doc_de)
                                        p_obj.alignment = WD_ALIGN_PARAGRAPH.CENTER

                                    if q_type == 3 and ep_q_copy.tag == qn('w:p'):
                                        ans = self.extract_and_clean_dapan(ep_q_copy)
                                        if ans: short_answer = ans
                                    
                                    t_texts_ep2 = []
                                    for t in ep_q_copy.iter(qn('w:t')):
                                        if t.text: t_texts_ep2.append(t.text)
                                    txt_ep = "".join(t_texts_ep2).strip().replace("\xa0", "").replace("\u200b", "")
                                    
                                    has_obj_ep = False
                                    for n in ep_q_copy.iter():
                                        if str(n.tag).endswith('imagedata') or str(n.tag).endswith('pic') or str(n.tag).endswith('object') or str(n.tag).endswith('drawing') or str(n.tag).endswith('oMath') or str(n.tag).endswith('AlternateContent'):
                                            has_obj_ep = True
                                            break
                                            
                                    if not txt_ep and not has_obj_ep and ep_q_copy.tag == qn('w:p'):
                                        continue 

                                    safe_append(body, ep_q_copy)
                                
                                if q_type == 3:
                                    p3_ans.append(short_answer)
                                    if ma == "GOC" and short_answer:
                                        p_ans_goc = OxmlElement('w:p')
                                        self.force_zero_spacing(p_ans_goc, justify=False, preserve_align=True)
                                        p_ans_obj = Paragraph(p_ans_goc, doc_de)
                                        run_ans = p_ans_obj.add_run(f"Đáp án: {short_answer}")
                                        run_ans.font.name = "Times New Roman"
                                        run_ans.font.size = Pt(12)
                                        run_ans.font.color.rgb = docx.shared.RGBColor(255, 0, 0)
                                        safe_append(body, p_ans_goc)
                                        
                                q_idx += 1
                                continue
                            
                            safe_append(body, p_q)
                            
                            q_elements_skip_first = []
                            for i in range(1, len(q['q_elements'])):
                                q_elements_skip_first.append(q['q_elements'][i])
                                
                            for el in q_elements_skip_first: 
                                ep_q_copy = copy.deepcopy(el)
                                if ep_q_copy.tag == qn('w:p'):
                                    ep_q_copy = self.unwrap_omathpara(ep_q_copy)
                                    ep_q_copy = self.clean_empty_lines(ep_q_copy)
                                    self.reset_paragraph_indent(ep_q_copy, left_cm=0, first_line_cm=0)
                                    self.force_zero_spacing(ep_q_copy, justify=False, preserve_align=True)
                                elif ep_q_copy.tag == qn('w:tbl'):
                                    ep_q_copy, q_idx = self.update_table_questions(ep_q_copy, q_idx)
                                    
                                t_texts_ep = []
                                for t in ep_q_copy.iter(qn('w:t')):
                                    if t.text: t_texts_ep.append(t.text)
                                txt_content = "".join(t_texts_ep).strip()
                                
                                has_img = False
                                for n in ep_q_copy.iter():
                                    if str(n.tag).endswith('imagedata') or str(n.tag).endswith('pic') or str(n.tag).endswith('object') or str(n.tag).endswith('drawing') or str(n.tag).endswith('AlternateContent'):
                                        has_img = True
                                        break
                                        
                                if has_img and not txt_content and ep_q_copy.tag == qn('w:p'):
                                    p_obj = Paragraph(ep_q_copy, doc_de)
                                    p_obj.alignment = WD_ALIGN_PARAGRAPH.CENTER

                                t_texts_ep2 = []
                                for t in ep_q_copy.iter(qn('w:t')):
                                    if t.text: t_texts_ep2.append(t.text)
                                txt_ep = "".join(t_texts_ep2).strip().replace("\xa0", "").replace("\u200b", "")
                                
                                has_obj_ep = False
                                for n in ep_q_copy.iter():
                                    if str(n.tag).endswith('imagedata') or str(n.tag).endswith('pic') or str(n.tag).endswith('object') or str(n.tag).endswith('drawing') or str(n.tag).endswith('oMath') or str(n.tag).endswith('AlternateContent'):
                                        has_obj_ep = True
                                        break
                                        
                                if not txt_ep and not has_obj_ep and ep_q_copy.tag == qn('w:p'):
                                    continue 

                                safe_append(body, ep_q_copy)
                            
                            ans_list = copy.deepcopy(q['answers'])
                            if ma != "GOC" and len(ans_list) >= 2 and do_shuffle_options:
                                fixed_positions = {}
                                for i, ans in enumerate(ans_list):
                                    if ans.get("is_fixed", False): fixed_positions[i] = ans
                                    
                                shuffleable_items = []
                                for i, ans in enumerate(ans_list):
                                    if not ans.get("is_fixed", False): shuffleable_items.append(ans)
                                    
                                random.shuffle(shuffleable_items)
                                
                                new_ans_list = []
                                shuf_idx = 0
                                for i in range(len(ans_list)):
                                    if i in fixed_positions:
                                        new_ans_list.append(fixed_positions[i])
                                    else:
                                        new_ans_list.append(shuffleable_items[shuf_idx])
                                        shuf_idx += 1
                                ans_list = new_ans_list
                                
                            labels = ['A', 'B', 'C', 'D', 'E', 'F', 'G']
                            if q_type == 2: labels = ['a', 'b', 'c', 'd', 'e', 'f', 'g']
                            
                            opts_ready = []
                            max_len_cm = 0
                            
                            q_p1_ans = ""
                            q_p2_ans = []
                            
                            for i, ans in enumerate(ans_list):
                                lbl = f"{labels[i]}. "
                                if q_type == 2: lbl = f"{labels[i]}) "
                                
                                blocks_list, l = self.get_option_data(ans, lbl)
                                
                                if blocks_list[0].tag == qn('w:p'):
                                    blocks_list[0] = self.clean_empty_lines(blocks_list[0])
                                    
                                if ma == "GOC" and ans.get("is_true", False):
                                    for r in blocks_list[0].xpath('.//w:r'):
                                        rPr = r.find(qn('w:rPr'))
                                        if rPr is None:
                                            rPr = OxmlElement('w:rPr')
                                            r.insert(0, rPr)
                                        color = rPr.find(qn('w:color'))
                                        if color is None:
                                            color = OxmlElement('w:color')
                                            rPr.append(color)
                                        color.set(qn('w:val'), 'FF0000')
                                        for attr in ('themeColor', 'themeTint', 'themeShade'):
                                            if qn(f'w:{attr}') in color.attrib:
                                                del color.attrib[qn(f'w:{attr}')]
                                    
                                opts_ready.append((blocks_list, l)) 
                                max_len_cm = max(max_len_cm, l)
                                
                                if q_type == 1 and ans.get("is_true", False): q_p1_ans = labels[i]
                                if q_type == 2:
                                    if ans.get("is_true", False): q_p2_ans.append("Đ")
                                    else: q_p2_ans.append("S")
                            
                            if q_type == 1: p1_ans.append(q_p1_ans)
                            if q_type == 2: p2_ans.append("".join(q_p2_ans))
                            
                            if q_type == 1:
                                if max_len_cm <= 4.3:
                                    self.append_options_as_text(doc_de, body, opts_ready, 4, 0.5)
                                elif max_len_cm <= 9.5:
                                    self.append_options_as_text(doc_de, body, opts_ready, 2, 0.5)
                                else:
                                    self.append_options_as_text(doc_de, body, opts_ready, 1, 0.5)
                            elif q_type == 2: 
                                for opt_tuple in opts_ready:
                                    blocks = opt_tuple[0]
                                    for b_idx, blk in enumerate(blocks):
                                        if blk.tag == qn('w:p'):
                                            self.force_zero_spacing(blk, justify=False, preserve_align=False)
                                            if b_idx == 0:
                                                self.force_part2_indent(blk, doc_de)
                                            Paragraph(blk, doc_de).alignment = WD_ALIGN_PARAGRAPH.LEFT
                                        safe_append(body, blk)
                            
                            q_idx += 1
                
                ps = []
                for p in body.xpath('.//w:p'): ps.append(p)
                for p in ps:
                    brs = []
                    for br in p.xpath('.//w:br'): brs.append(br)
                    for br in brs:
                        if br.get(qn('w:type')) == 'page':
                            br.getparent().remove(br)
                            
                    pgbs = []
                    for pgb in p.xpath('.//w:lastRenderedPageBreak'): pgbs.append(pgb)
                    for pgb in pgbs:
                        pgb.getparent().remove(pgb)

                ps2 = []
                for p in body.xpath('./w:p | .//w:tc/w:p'): ps2.append(p)
                for p in ps2:
                    t_texts = []
                    for t in p.xpath('.//w:t | .//m:t'):
                        if t.text: t_texts.append(t.text)
                    text_content = "".join(t_texts).strip().replace("\xa0", "").replace("\u200b", "")
                    
                    has_media = False
                    for n in p.iter():
                        if str(n.tag).endswith('imagedata') or str(n.tag).endswith('pic') or str(n.tag).endswith('object') or str(n.tag).endswith('drawing') or str(n.tag).endswith('oMath') or str(n.tag).endswith('AlternateContent'):
                            has_media = True
                            break
                            
                    in_table = False
                    if p.getparent() is not None and p.getparent().tag == qn('w:tc'): in_table = True
                    
                    if p.get(qn('w:rsidR')) == 'KEEP_ME_EMPTY_LINE':
                        continue
                        
                    if not text_content and not has_media:
                        if not in_table:
                            if p.getparent() is not None:
                                p.getparent().remove(p)
                    else:
                        if not in_table:
                            p_obj = Paragraph(p, doc_de)
                            old_space = p_obj.paragraph_format.space_before
                            
                            self.force_zero_spacing(p, justify=False, preserve_align=True)
                            
                            if old_space is not None and old_space > 0:
                                p_obj.paragraph_format.space_before = old_space
                            
                p_end_elm = OxmlElement('w:p')
                self.force_zero_spacing(p_end_elm, justify=False, preserve_align=True)
                p_end = Paragraph(p_end_elm, doc_de)
                p_end.alignment = WD_ALIGN_PARAGRAPH.CENTER
                
                end_text = "--------- THE END ---------" if is_english_exam else "--------- HẾT ---------"
                run_end = p_end.add_run(end_text)
                
                run_end.font.name = "Times New Roman"
                run_end.font.size = Pt(12)
                run_end.font.bold = True
                safe_append(body, p_end_elm)
                
                self.apply_global_font(doc_de)
                self.remove_youngmix_tags(doc_de)
                if ma != "GOC":
                    self.clear_answers_from_exam(doc_de)
                
                for p in doc_de.paragraphs:
                    old_space = p.paragraph_format.space_before
                    self.force_zero_spacing(p._element, justify=False, preserve_align=True)
                    if old_space is not None and old_space > 0:
                        p.paragraph_format.space_before = old_space
                        
                    if re.match(r'^\s*PHẦN\s+([IVX]+)', p.text, re.IGNORECASE):
                        p.paragraph_format.space_before = Cm(0.3)
                        p.paragraph_format.space_after = Cm(0.1)

                try:
                    if ma == "GOC":
                        doc_de.save(os.path.join(out_dir, f"De_Ma_GOC.docx"))
                    else:
                        doc_de.save(os.path.join(out_dir, f"De_Ma_{ma}.docx"))
                except PermissionError:
                    pass
                
                answers_data[ma]['P1'] = p1_ans
                answers_data[ma]['P2'] = p2_ans
                answers_data[ma]['P3'] = p3_ans

                if ma != "GOC":
                    if not is_first_for_all:
                        doc_all.add_section(docx.enum.section.WD_SECTION_START.NEW_PAGE)
                    
                    sect_all = doc_all.sections[-1]
                    sect_all.page_width = Cm(21.0)
                    sect_all.page_height = Cm(29.7)
                    sect_all.top_margin = Cm(1.5)
                    sect_all.bottom_margin = Cm(1.5)
                    sect_all.left_margin = Cm(1.5)
                    sect_all.right_margin = Cm(1.5)
                    sect_all.header_distance = Cm(0.8)
                    sect_all.footer_distance = Cm(0.6)
                    
                    sectPr_all = sect_all._sectPr
                    pgNumType = sectPr_all.find(qn('w:pgNumType'))
                    if pgNumType is None:
                        pgNumType = OxmlElement('w:pgNumType')
                        sectPr_all.append(pgNumType)
                    pgNumType.set(qn('w:start'), '1')

                    self._setup_footer(sect_all.footer, ma, "SECTIONPAGES", is_english_exam)
                    self._setup_footer(sect_all.first_page_footer, ma, "SECTIONPAGES", is_english_exam)

                    child_nodes = list(doc_de._body._body)
                    for child in child_nodes:
                        if child.tag != qn('w:sectPr'):
                            safe_append(body_all, child)
                           
                    is_first_for_all = False
            
            doc_all.add_section(docx.enum.section.WD_SECTION_START.NEW_PAGE)
            sect_hd = doc_all.sections[-1]
            sect_hd.page_width = Cm(21.0)
            sect_hd.page_height = Cm(29.7)
            sect_hd.top_margin = Cm(1.5)
            sect_hd.bottom_margin = Cm(1.5)
            sect_hd.left_margin = Cm(1.5)
            sect_hd.right_margin = Cm(1.5)
            sect_hd.header_distance = Cm(0.8)
            sect_hd.footer_distance = Cm(0.6)
            
            sect_hd.footer.is_linked_to_previous = False
            for p in sect_hd.footer.paragraphs:
                p.text = ""
            
            self.add_hd_cham_header(doc_all)
            
            p_space = doc_all.add_paragraph()
            self.force_zero_spacing(p_space._element, False, False)
            
            erws = []
            wrws = []
            
            hex_excel = ["Câu\\Mã đề", "GOC"]
            for m in ma_des_full:
                if m != "GOC": hex_excel.append(m)
            erws.append(hex_excel)
            
            hex_word = ["Câu\\Mã đề"]
            for m in ma_des_full:
                if m != "GOC": hex_word.append(m)
            wrws.append(hex_word)
            
            def clean_ans(val):
                return str(val).replace(',', '.').strip()
            
            excel_q_counter = 1
            parts_info = [(1, "PHẦN I", 'P1'), (2, "PHẦN II", 'P2'), (3, "PHẦN III", 'P3')]
            for pid, pt, pky in parts_info:
                mxps = [0]
                for m in ma_des_full:
                    if m in answers_data:
                        if pky in answers_data[m]:
                            mxps.append(len(answers_data[m][pky]))
                mxp = max(mxps)
                
                if mxp > 0:
                    empty_row = [pt]
                    for _ in range(len(hex_word) - 1): empty_row.append("")
                    wrws.append(empty_row)
                    
                    for i in range(mxp):
                        rex = [str(excel_q_counter)] 
                        rw = [str(i + 1)] 
                        
                        ag = []
                        if "GOC" in answers_data:
                            if pky in answers_data["GOC"]:
                                ag = answers_data["GOC"][pky]
                                
                        ans_goc = ""
                        if i < len(ag): ans_goc = ag[i]
                        rex.append(clean_ans(ans_goc))
                        
                        for m in ma_des_full:
                            if m == "GOC": continue
                            
                            am = []
                            if m in answers_data:
                                if pky in answers_data[m]:
                                    am = answers_data[m][pky]
                                    
                            val = ""
                            if i < len(am): val = am[i]
                            val = clean_ans(val)
                            
                            rex.append(val)
                            rw.append(val)
                            
                        erws.append(rex)
                        wrws.append(rw)
                        excel_q_counter += 1

            num_rows = len(wrws)
            num_cols = len(hex_word)
            if num_rows > 0 and num_cols > 0:
                tb_ans = doc_all.add_table(rows=num_rows, cols=num_cols)
                
                try:
                    tb_ans.style = 'Table Grid'
                except:
                    tblPr = tb_ans._element.find(qn('w:tblPr'))
                    if tblPr is not None:
                        borders = OxmlElement('w:tblBorders')
                        for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
                            border = OxmlElement(f'w:{side}')
                            border.set(qn('w:val'), 'single')
                            border.set(qn('w:sz'), '4')
                            border.set(qn('w:space'), '0')
                            border.set(qn('w:color'), 'auto')
                            borders.append(border)
                        tblPr.append(borders)

                tb_ans.alignment = WD_TABLE_ALIGNMENT.CENTER
                for r_idx, row_data in enumerate(wrws):
                    for c_idx, cell_data in enumerate(row_data):
                        cell = tb_ans.cell(r_idx, c_idx)
                        cell.text = str(cell_data)
                        
                        tcPr = cell._tc.get_or_add_tcPr()
                        vAlign = OxmlElement('w:vAlign')
                        vAlign.set(qn('w:val'), 'center')
                        tcPr.append(vAlign)
                        
                        for p in cell.paragraphs:
                            self.force_zero_spacing(p._element, justify=False, preserve_align=False)
                            self.reset_paragraph_indent(p._element, left_cm=0, first_line_cm=0) 
                            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                            for r in p.runs:
                                r.font.name = "Times New Roman"
                                r.font.size = Pt(12)
                            
                for r_idx, row_data in enumerate(wrws):
                    is_part = str(row_data[0]).startswith("PHẦN")
                    if is_part:
                        a = tb_ans.cell(r_idx, 0)
                        b = tb_ans.cell(r_idx, num_cols - 1)
                        a.merge(b)
                        a.text = str(row_data[0])
                        
                        tcPr = a._tc.get_or_add_tcPr()
                        vAlign = OxmlElement('w:vAlign')
                        vAlign.set(qn('w:val'), 'center')
                        tcPr.append(vAlign)
                        
                        for p in a.paragraphs:
                            self.force_zero_spacing(p._element, justify=False, preserve_align=False)
                            self.reset_paragraph_indent(p._element, left_cm=0, first_line_cm=0)
                            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                            for r in p.runs:
                                r.font.bold = True
                                r.font.name = "Times New Roman"
                                r.font.size = Pt(12)
                    elif r_idx == 0: 
                        for c_idx in range(num_cols):
                            c = tb_ans.cell(r_idx, c_idx)
                            for p in c.paragraphs:
                                self.force_zero_spacing(p._element, justify=False, preserve_align=False)
                                self.reset_paragraph_indent(p._element, left_cm=0, first_line_cm=0)
                                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                                for r in p.runs:
                                    r.font.bold = True
                                    r.font.name = "Times New Roman"
                                    r.font.size = Pt(12)
            
            try:
                for p in doc_all.paragraphs:
                    old_space = p.paragraph_format.space_before
                    self.force_zero_spacing(p._element, justify=False, preserve_align=True)
                    if old_space is not None and old_space > 0:
                        p.paragraph_format.space_before = old_space
                        
                    if re.match(r'^\s*PHẦN\s+([IVX]+)', p.text, re.IGNORECASE):
                        p.paragraph_format.space_before = Cm(0.3)
                        p.paragraph_format.space_after = Cm(0.1)
                        
                self.apply_global_font(doc_all)
                doc_all.save(os.path.join(out_dir, "Tat_Ca_Cac_De_Va_Dap_An.docx"))
            except PermissionError:
                raise Exception("Không thể lưu file Tat_Ca_Cac_De_Va_Dap_An.docx. Vui lòng đóng file nếu đang mở.")
            
            if HAS_OPENPYXL:
                wb = openpyxl.Workbook()
                ws = wb.active
                ws.title = "Đáp Án TNMaker"
                for r in erws:
                    ws.append(r)
                for col in ws.columns:
                    for cell in col:
                        cell.alignment = Alignment(horizontal='center', vertical='center')
                        if cell.row == 1 or cell.value in ("PHẦN I", "PHẦN II", "PHẦN III"):
                            cell.font = Font(name="Times New Roman", bold=True, size=12)
                        else:
                            cell.font = Font(name="Times New Roman", size=12)
                ws.column_dimensions['A'].width = 15
                for i in range(len(hex_excel)-1):
                    ws.column_dimensions[get_column_letter(i+2)].width = 15
                excel_path = os.path.join(out_dir, "Dap an dinh dang SmartTest 2025.xlsx")
                try:
                    wb.save(excel_path)
                except PermissionError:
                    pass
            else:
                excel_path = os.path.join(out_dir, "Dap an dinh dang SmartTest 2025.csv")
                with open(excel_path, "w", newline="", encoding="utf-8-sig") as f:
                    csv.writer(f).writerows(erws)

        except Exception as e: 
            raise e
        finally:
            gc.collect()